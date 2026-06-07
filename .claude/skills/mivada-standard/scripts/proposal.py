#!/usr/bin/env python3
"""Mivada Standard — native A4 proposal document (.docx).

A real, editable Word proposal in the design language (Inter, coral accent, ink
headings, coral-header tables, black callout bands) — not slide screenshots.
Build the sample, or import `Proposal` to compose your own sections.

  python proposal.py                 # -> Mivada_Proposal.docx (sample AMS proposal)
"""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CORAL = RGBColor(0xEA, 0x49, 0x3F); INK = RGBColor(0x11, 0x11, 0x11)
GRAPH = RGBColor(0x32, 0x32, 0x32); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MID = RGBColor(0xAE, 0xAE, 0xAE)
HAIR = "E4E4E0"; BLACK = "111111"; CORALHEX = "EA493F"; FONT = "Inter"
HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "..", "assets", "logos", "Mivada_Logo_Master_RGB_L.png")
CONTENT_MM = 170  # A4 210 - 2*20 margins

def _el(tag): return OxmlElement(tag)
# schema child order for tcPr / tblPr — insert respecting these so the .docx validates
_TCPR = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
         "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
_TBLPR = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
          "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
          "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"]
def _ordered_insert(parent, el, order):
    tag = "w:" + el.tag.split("}")[-1]
    old = parent.find(qn(tag))
    if old is not None: parent.remove(old)
    idx = order.index(tag)
    for child in parent:
        ct = "w:" + child.tag.split("}")[-1]
        if ct in order and order.index(ct) > idx:
            child.addprevious(el); return
    parent.append(el)
def _shade(cell, fill):
    sh = _el("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill)
    _ordered_insert(cell._tc.get_or_add_tcPr(), sh, _TCPR)
_SIDE = {"top": 0, "start": 1, "bottom": 2, "end": 3}
def _cell_borders(cell, color=HAIR, sz=4, sides=("top", "start", "bottom", "end")):
    b = _el("w:tcBorders")
    for s in sorted(sides, key=lambda x: _SIDE[x]):
        e = _el("w:" + s); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color); e.set(qn("w:space"), "0"); b.append(e)
    _ordered_insert(cell._tc.get_or_add_tcPr(), b, _TCPR)
def _cell_margins(cell, t=70, b=70, l=120, r=120):
    m = _el("w:tcMar")
    for k, v in (("top", t), ("start", l), ("bottom", b), ("end", r)):
        e = _el("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    _ordered_insert(cell._tc.get_or_add_tcPr(), m, _TCPR)
_PPR_AFTER = ["w:shd", "w:tabs", "w:spacing", "w:ind", "w:contextualSpacing", "w:jc",
              "w:textAlignment", "w:outlineLvl", "w:rPr", "w:sectPr"]
def _para_border(p, side="bottom", color=HAIR, sz=6, space=2):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("w:pBdr"))
    if old is not None: pPr.remove(old)
    bdr = _el("w:pBdr"); e = _el("w:" + side)
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), str(space)); e.set(qn("w:color"), color)
    bdr.append(e)
    ref = next((pPr.find(qn(t)) for t in _PPR_AFTER if pPr.find(qn(t)) is not None), None)
    ref.addprevious(bdr) if ref is not None else pPr.append(bdr)
def _add_page_field(p, color, size):
    r = p.add_run(); r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
    b = _el("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = _el("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " PAGE "
    e = _el("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._r.append(b); r._r.append(i); r._r.append(e)
def _no_cell_borders(cell):
    b = _el("w:tcBorders")
    for s in ("top", "start", "bottom", "end"):
        e = _el("w:" + s); e.set(qn("w:val"), "nil"); b.append(e)
    _ordered_insert(cell._tc.get_or_add_tcPr(), b, _TCPR)

class Proposal:
    def __init__(self):
        self.doc = Document()
        n = self.doc.styles["Normal"]; n.font.name = FONT; n.font.size = Pt(10.5); n.font.color.rgb = GRAPH
        s = self.doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = Mm(20); s.left_margin = s.right_margin = Mm(20)
        s.header_distance = Mm(12); s.footer_distance = Mm(12)
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None: z.set(qn("w:percent"), "100")
        self._footer(s)

    def run(self, p, text, size=10.5, bold=False, color=GRAPH, caps=False):
        r = p.add_run(text.upper() if caps else text)
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        return r
    def _para(self, before=0, after=6, line=None, align=None):
        p = self.doc.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        if line: pf.line_spacing = line
        if align: p.alignment = align
        return p

    def eyebrow(self, text, before=14):
        p = self._para(before, 2); self.run(p, text, 8.5, True, CORAL, caps=True)
    def headline(self, ink, coral="", before=2, size=18):
        p = self._para(before, 8); self.run(p, ink, size, True, INK)
        if coral: self.run(p, coral, size, True, CORAL)
    def section(self, eyebrow, ink, coral=""):
        self.eyebrow(eyebrow); self.headline(ink, coral)
    def body(self, text, after=8):
        p = self._para(0, after, line=1.3); self.run(p, text, 10.5, False, GRAPH)
    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
            self.run(p, it, 10.5, False, GRAPH)
    def rule(self, before=2, after=8):
        p = self._para(before, after); _para_border(p)

    def _table(self, widths_mm):
        t = self.doc.add_table(rows=0, cols=len(widths_mm)); t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        w = _el("w:tblW"); w.set(qn("w:w"), str(int(CONTENT_MM * 56.69))); w.set(qn("w:type"), "dxa")
        _ordered_insert(t._tbl.tblPr, w, _TBLPR)
        return t

    def kpis(self, items):
        """items=[(num, label)] — big coral numbers in a borderless row."""
        t = self._table([CONTENT_MM / len(items)] * len(items)); row = t.add_row()
        for i, (num, label) in enumerate(items):
            c = row.cells[i]; c.width = Mm(CONTENT_MM / len(items)); _no_cell_borders(c); _cell_margins(c, 60, 60, 0, 120)
            if i: _cell_borders(c, HAIR, 6, sides=("start",))
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            self._cell_run(c.paragraphs[0], num, 26, True, CORAL)
            p2 = c.add_paragraph(); self._cell_run(p2, label, 8.5, True, MID, caps=True)
        self._para(0, 6)

    def callout(self, lead, rest):
        t = self._table([CONTENT_MM]); c = t.add_row().cells[0]; c.width = Mm(CONTENT_MM)
        _shade(c, BLACK); _no_cell_borders(c); _cell_margins(c, 140, 140, 200, 200)
        p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.25
        self._cell_run(p, lead + "  ", 11, True, CORAL); self._cell_run(p, rest, 11, False, WHITE)
        self._para(8, 8)

    def table(self, cols, rows):
        """cols=[(name,width_mm)]; rows=[[cell,…]]. Coral header, bordered body."""
        t = self._table([w for _, w in cols]); hdr = t.add_row()
        for i, (name, wmm) in enumerate(cols):
            c = hdr.cells[i]; c.width = Mm(wmm); _shade(c, CORALHEX); _no_cell_borders(c); _cell_margins(c)
            self._cell_run(c.paragraphs[0], name, 9, True, WHITE, caps=True)
        for r_i, row in enumerate(rows):
            tr = t.add_row()
            for i, (val) in enumerate(row):
                c = tr.cells[i]; c.width = Mm(cols[i][1]); _cell_borders(c, HAIR, 4); _cell_margins(c)
                strong = (i == 0)
                self._cell_run(c.paragraphs[0], val, 9.5, strong, INK if strong else GRAPH)
        self._para(8, 8)

    def _cell_run(self, p, text, size, bold, color, caps=False):
        r = p.add_run(text.upper() if caps else text)
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1

    def masthead(self, meta):
        p = self._para(0, 6); p.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), WD_TAB_ALIGNMENT.RIGHT)
        if os.path.exists(LOGO):
            p.add_run().add_picture(LOGO, width=Mm(38))
        else:
            self.run(p, "Mivada", 15, True, INK)
        self.run(p, "\t" + meta, 8.5, True, GRAPH, caps=True)
        _para_border(p, "bottom", BLACK, 16, 6)

    def lede(self, eyebrow, ink, coral):
        self.eyebrow(eyebrow, before=22)
        p = self._para(2, 12); self.run(p, ink, 26, True, INK); self.run(p, coral, 26, True, CORAL)

    def pagebreak(self):
        self.doc.add_page_break()

    def _footer(self, section):
        f = section.footer.paragraphs[0]
        f.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), WD_TAB_ALIGNMENT.RIGHT)
        _para_border(f, "top", HAIR, 6, 4)
        self.run(f, "mivada.com", 8, True, GRAPH)
        self.run(f, "    ·    Commercial in confidence", 8, False, MID)
        self.run(f, "\t", 8, False, MID)
        _add_page_field(f, MID, 8)

    def save(self, path): self.doc.save(path); return path


def sample():
    """A generic Workday Managed Services (AMS) proposal."""
    p = Proposal()
    p.masthead("Workday Managed Services · Proposal · 2026")
    p.lede("Technology, human first", "Workday support that ", "owns the outcome.")
    p.body("This proposal sets out how Mivada runs and continuously improves your Workday platform — a "
           "blended onshore–offshore managed service, governed on ITIL, measured against clear SLAs, and "
           "delivered by certified consultants who treat your outcomes as their own.")
    p.kpis([("2014", "Founded"), ("200+", "Specialists · AU, US & India"), ("24×7", "P1/P2 on-call")])

    p.section("Our approach", "Across the whole Workday ", "lifecycle.")
    p.body("We support every stage — advise, implement, optimise and manage — with one accountable team, "
           "so knowledge compounds instead of being handed off.")
    p.bullets([
        "Advise — roadmap, business case and readiness, before any build.",
        "Implement & optimise — certified consultants leading to Workday best practice.",
        "Manage — Application Managed Services and Managed Payroll, 24×7, with continuous improvement.",
    ])

    p.section("Coverage model", "Seamless onshore–offshore ", "coverage.")
    p.body("Onshore in Australia owns governance, P1/P2 resolution and start-of-day health checks; offshore "
           "in India runs day-to-day incidents, requests and monitoring. 24×7 on-call is the always-on safety net.")

    p.section("Service levels", "Measurable ", "service levels.")
    p.table([("Priority", 26), ("Definition", 86), ("Acknowledge", 28), ("Resolution", 30)], [
        ["P1", "Critical — a vital business process cannot complete (e.g. payroll).", "30 min", "4 hours"],
        ["P2", "Critical severity with moderate impact, or urgent major-business impact.", "1 hour", "8 hours"],
        ["P3", "Moderate impact where a workaround is known and available.", "4 hours", "3 days"],
        ["P4", "Low impact and/or low urgency; no significant business impact.", "1 day", "5 days"],
    ])

    p.section("Commercials", "Engagement ", "pricing.")
    p.table([("Item", 80), ("Detail", 90)], [
        ["Allocated capacity", "100 hours / month"],
        ["Monthly fee (AUD)", "$15,000 — predictable managed service"],
        ["Annual commitment (AUD)", "$180,000"],
        ["Blended rate (AUD)", "$150 / hour · 40 / 60 onshore–offshore"],
    ])
    p.body("Up to 20% of unused monthly hours carry forward; hours flex with demand and are reviewed against "
           "actual utilisation every six months. Figures are illustrative pending scope confirmation.", after=8)

    p.callout("One team. One platform. Always on.", "Workday GO's preconfigured tenant removes design "
              "overhead — we validate, tailor and run, so your people feel the system working from day one.")

    p.section("Next steps", "Let's build something ", "human.")
    p.bullets([
        "Confirm scope, modules in support, and the coverage window.",
        "Agree SLAs, governance cadences and the transition plan.",
        "Mobilise — governance live in week one, leading AMS by week four.",
    ])
    return p.save(os.path.join(os.getcwd(), "Mivada_Proposal.docx"))


if __name__ == "__main__":
    print("wrote", sample())
