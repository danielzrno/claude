#!/usr/bin/env python3
# ============================================================================
# Mivada Editorial — STANDARD SLIDE LIBRARY
# Content-driven, brand-styled builders for the key slide types. Each takes a
# Deck `d` plus content and adds one slide. Compose a deck from these + the
# Deck methods in build_pptx.py (cover_plain, kpis, pillars, steps, split,
# pullquote, reasons, contact_plain, content, quote_plain, chart_coral).
# ============================================================================
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_pptx import Deck, THEME, PAGE_W
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

CORAL, INK, WHITE, DARK = "EA493F", "111111", "FFFFFF", "141414"
M = 0.96; CW = PAGE_W - 2 * M
LOGO_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "logos")

# ---------------- shared primitives -----------------------------------------
def head(d, s, pre, accent, tag, dark=False, y=1.0, hsize=33):
    """Title (ink/coral) + single coral section slug, top-right."""
    d._slug(s, tag, dark=dark)
    _, tf = d._box(s, 0.92, y, CW, 1.1); p = d._para(tf, first=True, line=1.0)
    d._run(p, pre + " ", hsize, (WHITE if dark else INK), bold=True, spacing=-0.022)
    if accent: d._run(p, accent, hsize, CORAL, bold=True, spacing=-0.022)

def standfirst(d, s, text, y=1.9, w=11.2, size=13, dark=False):
    _, tf = d._box(s, M, y, w, 0.6)
    d._run(d._para(tf, first=True, line=1.25), text, size, THEME["rev_soft"] if dark else "323232", bold=False, spacing=0)

def bullets(d, s, x, y, w, items, size=11, dark=False, gap=0.42):
    cy = y
    for it in items:
        _, tf = d._box(s, x, cy, w, gap); p = d._para(tf, first=True, line=1.12)
        d._run(p, "—  ", size, CORAL, bold=True, spacing=0)
        d._run(p, it, size, (THEME["rev_body"] if dark else "323232"), bold=False, spacing=0)
        cy += gap

def tile(d, s, x, y, w, h, text, fill=WHITE, line=None, tc=INK, size=9.5):
    d._rect(s, x, y, w, h, fill_hex=fill, line_hex=line, line_w=1.0)
    _, tf = d._box(s, x + 0.08, y, w - 0.16, h, anchor=MSO_ANCHOR.MIDDLE)
    p = d._para(tf, first=True, align=PP_ALIGN.CENTER, line=1.0)
    d._run(p, text, size, tc, bold=True, spacing=-0.005)

def band(d, s, x, y, w, h, label, items=None, fill=INK, tc=WHITE):
    d._rect(s, x, y, w, h, fill_hex=fill)
    _, tf = d._box(s, x + 0.25, y, w - 0.5, h, anchor=MSO_ANCHOR.MIDDLE); p = d._para(tf, first=True)
    d._run(p, label.upper(), 11.5, tc, bold=True, spacing=0.1, caps=True)
    if items: d._run(p, "    " + items, 10, tc, bold=False, spacing=0)

def _card(d, s, x, y, w, h, label, sub, items, dark=True):
    d._rect(s, x, y, w, h, fill_hex=(DARK if dark else WHITE),
            line_hex=(THEME["rev_hair"] if dark else THEME["hairline"]), line_w=1.0)
    _, tf = d._box(s, x + 0.3, y + 0.26, w - 0.55, 0.4)
    d._run(d._para(tf, first=True), label.upper(), 12.5, CORAL, bold=True, spacing=0.12, caps=True)
    _, tf = d._box(s, x + 0.3, y + 0.66, w - 0.55, 0.5)
    d._run(d._para(tf, first=True, line=1.05), sub, 14.5, (WHITE if dark else INK), bold=True, spacing=-0.01)
    bullets(d, s, x + 0.3, y + 1.32, w - 0.6, items, size=10.5, dark=dark, gap=0.46)

# ============================================================================
# STANDARD SLIDES
# ============================================================================
def logo_slide(d, pre="Technology,", accent="human first.", logo="Mivada_Logo_Master_RGB_L"):
    """White brand-arrival slide: centred master logo + a one-line tagline."""
    s = d._slide(WHITE)
    p = os.path.join(LOGO_DIR, logo + ".png")
    if os.path.exists(p):
        from PIL import Image as I
        iw, ih = I.open(p).size; h = 0.98; w = h * iw / ih
        s.shapes.add_picture(p, Inches((PAGE_W - w) / 2), Inches(2.72), height=Inches(h))
    _, tf = d._box(s, 2, 4.12, PAGE_W - 4, 0.6); pr = d._para(tf, first=True, align=PP_ALIGN.CENTER)
    d._run(pr, pre + " ", 16, INK, bold=True, spacing=0); d._run(pr, accent, 16, CORAL, bold=True, spacing=0)
    return s

def divider(d, eyebrow, pre, accent, sub=None):
    """Black section divider."""
    s = d._slide(THEME["black"]); d._slug(s, eyebrow, dark=True)
    _, tf = d._box(s, 0.92, 2.95, 11.4, 1.6); p = d._para(tf, first=True, line=0.98)
    d._run(p, pre + " ", 52, WHITE, bold=True, spacing=-0.025); d._run(p, accent, 52, CORAL, bold=True, spacing=-0.025)
    if sub:
        _, tf = d._box(s, 0.96, 4.55, 10.5, 0.6); d._run(d._para(tf, first=True), sub, 15, THEME["rev_body"], bold=False, spacing=0)
    return s

def statement(d, pre, accent, tag, sub=None, points=None, footer=None, dark=True):
    """Manifesto: big two-tone statement + optional 2-col points + footer line."""
    s = d._slide(THEME["black"] if dark else THEME["off_white"]); d._slug(s, tag, dark=dark)
    _, tf = d._box(s, 0.92, 1.36, 11.6, 1.8); p = d._para(tf, first=True, line=0.98)
    d._run(p, pre, 38, (WHITE if dark else INK), bold=True, spacing=-0.025)
    if accent:
        p2 = d._para(tf, line=0.98); d._run(p2, accent, 38, CORAL, bold=True, spacing=-0.025)
    yy = 3.4
    if sub:
        _, tf = d._box(s, 0.96, 3.12, 11, 0.5); d._run(d._para(tf, first=True), sub, 15, THEME["rev_body"] if dark else "323232", bold=True, spacing=0); yy = 3.95
    if points:
        per = (len(points) + 1) // 2
        for ci, group in enumerate((points[:per], points[per:])):
            cx = M + ci * (CW / 2)
            for ri, (key, val) in enumerate(group):
                ry = yy + ri * 0.6
                _, tf = d._box(s, cx, ry, CW / 2 - 0.4, 0.5); pp = d._para(tf, first=True)
                d._run(pp, "—   ", 14, CORAL, bold=True, spacing=0)
                d._run(pp, key + ("    " if val else ""), 15, (WHITE if dark else INK), bold=True, spacing=-0.01)
                if val: d._run(pp, val, 12.5, THEME["rev_soft"] if dark else "323232", bold=False, spacing=0)
    if footer:
        d._hline(s, M, 6.74, CW, THEME["rev_hair"] if dark else THEME["hairline"], 1.0)
        _, tf = d._box(s, M, 6.88, CW, 0.4); pf = d._para(tf, first=True)
        d._run(pf, footer, 11, CORAL, bold=True, spacing=0.04)
    return s

def logo_wall(d, pre, accent, tag, names=None, image=None):
    """Client wall — a colour-logo image if given, else a clean name grid."""
    if image and os.path.exists(image):
        s = d._slide(WHITE); d._slug(s, tag)
        _, tf = d._box(s, 0.92, 1.0, CW, 1.0); p = d._para(tf, first=True)
        d._run(p, pre + " ", 33, INK, bold=True, spacing=-0.022); d._run(p, accent, 33, CORAL, bold=True, spacing=-0.022)
        from PIL import Image as I
        iw, ih = I.open(image).size; w = 10.8; h = w * ih / iw
        if h > 3.9: h = 3.9; w = h * iw / ih
        s.shapes.add_picture(image, Inches((PAGE_W - w) / 2), Inches(2.7), width=Inches(w)); return s
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    standfirst(d, s, "A decade of work with some of Australia's most demanding operations.", 2.05)
    cols = 4; cw = CW / cols; y0 = 2.95; rowh = 0.62
    for i, nm in enumerate(names or []):
        r, c = divmod(i, cols); x = M + c * cw; y = y0 + r * rowh
        _, tf = d._box(s, x, y, cw - 0.2, 0.5); d._run(d._para(tf, first=True), nm, 15, INK, bold=True, spacing=-0.01)
        d._hline(s, x, y + 0.5, cw - 0.35, THEME["hairline"], 1.0)
    return s

def cases(d, pre, accent, tag, items, sub=None):
    """3-column case studies (client + paragraph)."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub, 2.0, size=13)
    col = CW / max(len(items), 1)
    for i, (client, para) in enumerate(items):
        cx = M + i * col
        if i > 0: d._vline(s, cx, 2.95, 3.8, THEME["hairline"], 1.0)
        pad = 0.0 if i == 0 else 0.32
        _, tf = d._box(s, cx + pad, 2.9, col - pad - 0.25, 0.5); d._run(d._para(tf, first=True), client.upper(), 12, CORAL, bold=True, spacing=0.06, caps=True)
        _, tf = d._box(s, cx + pad, 3.45, col - pad - 0.28, 3.4); d._run(d._para(tf, first=True, line=1.28), para, 10.5, "323232", bold=False, spacing=0)
    return s

def kpi_grid(d, pre, accent, tag, stats, sub=None):
    """Capability numbers — up to 7 stats over two rows. stats=[(num,suffix,label)]."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub, 2.05)
    top = stats[:4] if len(stats) > 4 else stats
    d._stat_row(s, 3.6, top, light=True)
    if len(stats) > 4: d._stat_row(s, 5.6, stats[4:], light=True)
    return s

def two_cards(d, pre, accent, tag, cards, sub=None, footer=None, dark=True):
    """Two comparison cards (label + sub + bullets)."""
    s = d._slide(THEME["black"] if dark else THEME["off_white"]); head(d, s, pre, accent, tag, dark=dark)
    if sub: standfirst(d, s, sub, dark=dark)
    cw = (CW - 0.4) / 2
    for i, (label, csub, items) in enumerate(cards[:2]):
        _card(d, s, M + i * (cw + 0.4), 2.55, cw, 3.6, label, csub, items, dark=dark)
    if footer:
        _, tf = d._box(s, M, 6.35, CW, 0.5)
        d._run(d._para(tf, first=True), footer, 12, THEME["mid_grey"], bold=True, spacing=0)
    return s

def columns(d, pre, accent, tag, cols, sub=None, footer=None):
    """2–3 columns of labelled lists; long lists auto-split into two sub-columns."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub)
    n = len(cols); colw = CW / n; ytop = 2.7
    for i, (heading, items) in enumerate(cols):
        cx = M + i * colw
        if i > 0: d._vline(s, cx, ytop, 3.5, THEME["hairline"], 1.0)
        pad = 0.0 if i == 0 else 0.34
        _, tf = d._box(s, cx + pad, ytop, colw - pad - 0.3, 0.5); d._run(d._para(tf, first=True), heading, 20, INK, bold=True, spacing=-0.02)
        if len(items) > 7:
            half = (len(items) + 1) // 2; subw = (colw - pad - 0.3) / 2
            bullets(d, s, cx + pad, ytop + 0.65, subw, items[:half], size=10.5, gap=0.38)
            bullets(d, s, cx + pad + subw, ytop + 0.65, subw, items[half:], size=10.5, gap=0.38)
        else:
            bullets(d, s, cx + pad, ytop + 0.65, colw - pad - 0.3, items, size=11.5, gap=0.5)
    if footer:
        _, tf = d._box(s, M, 6.6, CW, 0.4); d._run(d._para(tf, first=True), footer, 12.5, CORAL, bold=True, spacing=0.02)
    return s

def framework_stack(d, pre, accent, tag, top_label, gov_label, gov_items, tiles, mid_label, foundations, sub=None):
    """Layered framework: top coral band → governance band → tile grid → mid band → 3 foundations."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub)
    fx, fw = M, CW
    band(d, s, fx, 2.5, fw, 0.34, top_label, fill=CORAL)
    band(d, s, fx, 2.88, fw, 0.46, gov_label, gov_items, fill=INK)
    gx, gy = 0.14, 0.14; cols = 4; tw = (fw - gx * (cols - 1)) / cols; th = 0.5; y0 = 3.5
    for i, t in enumerate(tiles):
        r, c = divmod(i, cols); x = fx + c * (tw + gx); y = y0 + r * (th + gy)
        tile(d, s, x, y, tw, th, t, fill=WHITE, line=THEME["hairline"], tc=INK, size=9.5)
    rows = (len(tiles) + cols - 1) // cols
    yb = y0 + rows * (th + gy)
    band(d, s, fx, yb, fw, 0.42, mid_label, fill=CORAL)
    fy = yb + 0.52; fwidth = (fw - 0.3) / 3
    for i, f in enumerate(foundations[:3]):
        x = fx + i * (fwidth + 0.15)
        tile(d, s, x, fy, fwidth, 0.4, f.upper(), fill="F2F3EE", line=THEME["hairline"], tc=CORAL, size=10)
    return s

def timeline_gantt(d, pre, accent, tag, hours_top, hours_bottom, bars, overlap, callout, descs, sub=None, axis="← AEST   ·   IST →"):
    """Two-track gantt over an hour grid. bars=[(label,start,end,fill)], overlap=(s,e)."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub)
    n = len(hours_top); gx0 = M; gw = CW; cw = gw / n; gy = 2.95
    for i, h in enumerate(hours_top):
        _, tf = d._box(s, gx0 + i * cw, gy - 0.28, cw, 0.24); d._run(d._para(tf, first=True, align=PP_ALIGN.CENTER), h, 8, THEME["mid_grey"], bold=True, spacing=0)
    for i in range(n + 1): d._vline(s, gx0 + i * cw, gy, 1.5, THEME["hairline"], 0.75)
    d._hline(s, gx0, gy, gw, THEME["hairline"], 0.75); d._hline(s, gx0, gy + 1.5, gw, THEME["hairline"], 0.75)
    if overlap: d._rect(s, gx0 + overlap[0] * cw, gy, (overlap[1] - overlap[0]) * cw, 1.5, fill_hex="FBE7E4", line_hex=None)
    for j, (label, a, b, fill) in enumerate(bars):
        by = gy + 0.18 + j * 0.68
        d._rect(s, gx0 + a * cw + 0.04, by, (b - a) * cw - 0.08, 0.46, fill_hex=fill)
        _, tf = d._box(s, gx0 + a * cw + 0.12, by + 0.04, 5, 0.36); d._run(d._para(tf, first=True), label, 11, WHITE, bold=True, spacing=0.02)
    for i, h in enumerate(hours_bottom):
        _, tf = d._box(s, gx0 + i * cw, gy + 1.54, cw, 0.24); d._run(d._para(tf, first=True, align=PP_ALIGN.CENTER), h, 8, THEME["mid_grey"], bold=True, spacing=0)
    _, tf = d._box(s, gx0, gy + 1.78, 4, 0.22); d._run(d._para(tf, first=True), axis, 8.5, CORAL, bold=True, spacing=0.06, caps=True)
    if descs:
        bullets(d, s, M, 5.15, 5.3, descs[0], size=10.5, gap=0.36)
        if len(descs) > 1: bullets(d, s, 6.5, 5.15, 4.0, descs[1], size=10.5, gap=0.36)
    if callout:
        d._rect(s, M, 6.5, CW, 0.46, fill_hex=CORAL)
        _, tf = d._box(s, M + 0.25, 6.57, CW - 0.5, 0.34); p = d._para(tf, first=True)
        d._run(p, callout[0] + "   ", 11, WHITE, bold=True, spacing=0.12, caps=True)
        d._run(p, callout[1], 11, WHITE, bold=False, spacing=0)
    return s

def pyramid_tiers(d, pre, accent, tag, tiers, footer=None, headers=("Customer", "Activities", "Mivada", "Cadence"), sub=None):
    """Central stepped pyramid flanked by Customer · Activities / Mivada · Cadence. tiers=[(name,cad,cust,act,miv)]."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub)
    for hx, w, htxt in [(M, 2.0, headers[0]), (3.05, 2.0, headers[1]), (8.3, 2.0, headers[2]), (10.4, 2.0, headers[3])]:
        _, tf = d._box(s, hx, 2.4, w, 0.3); d._run(d._para(tf, first=True), htxt.upper(), 9, THEME["mid_grey"], bold=True, spacing=0.12, caps=True)
    y0 = 2.78; bh = 1.18; pyr_cx = PAGE_W / 2; pyr_w = [1.5, 2.3, 3.1]; pyr_fill = [CORAL, "D8392E", INK]
    for i, (name, cad, cust, act, miv) in enumerate(tiers[:3]):
        y = y0 + i * bh; w = pyr_w[i]
        d._rect(s, pyr_cx - w / 2, y + 0.12, w, bh - 0.26, fill_hex=pyr_fill[i])
        _, tf = d._box(s, pyr_cx - w / 2, y + 0.12, w, bh - 0.26, anchor=MSO_ANCHOR.MIDDLE)
        d._run(d._para(tf, first=True, align=PP_ALIGN.CENTER), name.upper(), 12, WHITE, bold=True, spacing=0.08, caps=True)
        _, tf = d._box(s, M, y + 0.05, 2.0, bh - 0.1); d._run(d._para(tf, first=True, line=1.16), cust, 9.5, "323232", bold=False, spacing=0)
        _, tf = d._box(s, 3.05, y + 0.05, 1.95, bh - 0.1); d._run(d._para(tf, first=True, line=1.16), act, 9.5, "323232", bold=False, spacing=0)
        _, tf = d._box(s, 8.3, y + 0.05, 2.0, bh - 0.1); d._run(d._para(tf, first=True, line=1.16), miv, 9.5, INK, bold=True, spacing=0)
        _, tf = d._box(s, 10.4, y + 0.16, 2.0, 0.4); d._run(d._para(tf, first=True), cad.upper(), 9.5, CORAL, bold=True, spacing=0.08, caps=True)
        d._hline(s, M, y, CW, THEME["hairline"], 0.75)
    d._hline(s, M, y0 + 3 * bh, CW, THEME["hairline"], 0.75)
    if footer:
        _, tf = d._box(s, M, y0 + 3 * bh + 0.12, CW, 0.4); d._run(d._para(tf, first=True), footer, 11, CORAL, bold=True, spacing=0.01)
    return s

def org_chart(d, pre, accent, tag, leaders, delivery, banners, sub=None):
    """Leadership row → delivery row (connected) + two framing banners."""
    s = d._slide(THEME["off_white"]); head(d, s, pre, accent, tag)
    if sub: standfirst(d, s, sub)
    cwL = (CW - 0.5) / 2
    def node(x, y, w, h, title, desc, fill=WHITE, tc=INK, dc="323232"):
        d._rect(s, x, y, w, h, fill_hex=fill, line_hex=(None if fill != WHITE else THEME["hairline"]), line_w=1.0)
        _, tf = d._box(s, x + 0.28, y + 0.2, w - 0.5, 0.4); d._run(d._para(tf, first=True), title, 15, tc, bold=True, spacing=-0.01)
        _, tf = d._box(s, x + 0.28, y + 0.62, w - 0.5, 0.6); d._run(d._para(tf, first=True, line=1.18), desc, 10.5, dc, bold=False, spacing=0)
    for i, (title, desc) in enumerate(leaders[:2]):
        node(M + i * (cwL + 0.5), 2.5, cwL, 1.05, title, desc)
    d._vline(s, M + cwL / 2, 3.55, 0.35, THEME["mid_grey"], 1.0)
    d._vline(s, M + cwL + 0.5 + cwL / 2, 3.55, 0.35, THEME["mid_grey"], 1.0)
    d._hline(s, M + cwL / 2, 3.9, cwL + 0.5, THEME["mid_grey"], 1.0)
    d._vline(s, PAGE_W / 2, 3.9, 0.3, THEME["mid_grey"], 1.0)
    for i, (title, desc) in enumerate(delivery[:2]):
        node(M + i * (cwL + 0.5), 4.2, cwL, 1.15, title, desc)
    bw = (CW - 0.5) / 2
    for i, (label, desc, fill) in enumerate(banners[:2]):
        x = M + i * (bw + 0.5); d._rect(s, x, 5.65, bw, 1.0, fill_hex=fill)
        tc = WHITE if fill != INK else CORAL; dc = WHITE if fill != INK else THEME["rev_body"]
        _, tf = d._box(s, x + 0.28, 5.78, bw - 0.5, 0.4); d._run(d._para(tf, first=True), label.upper(), 11, tc, bold=True, spacing=0.1, caps=True)
        _, tf = d._box(s, x + 0.28, 6.14, bw - 0.5, 0.45); d._run(d._para(tf, first=True, line=1.12), desc, 10.5, dc, bold=True, spacing=0)
    return s


def cards3(d, pre, accent, tag, cards, sub=None, footer=None, dark=True):
    """Three boxes, each a label + a short paragraph. Less text; for 'what / why / how'."""
    s = d._slide(THEME["black"] if dark else THEME["off_white"]); head(d, s, pre, accent, tag, dark=dark)
    if sub: standfirst(d, s, sub, dark=dark)
    bw = (CW - 0.8) / 3; y = 2.75; h = 3.3
    for i, (label, body) in enumerate(cards[:3]):
        x = M + i * (bw + 0.4)
        d._rect(s, x, y, bw, h, fill_hex=(DARK if dark else WHITE),
                line_hex=(THEME["rev_hair"] if dark else THEME["hairline"]), line_w=1.0)
        _, tf = d._box(s, x + 0.3, y + 0.32, bw - 0.55, 0.4)
        d._run(d._para(tf, first=True), label.upper(), 12.5, CORAL, bold=True, spacing=0.1, caps=True)
        _, tf = d._box(s, x + 0.3, y + 0.92, bw - 0.6, h - 1.2)
        d._run(d._para(tf, first=True, line=1.4), body, 13, (THEME["rev_body"] if dark else "323232"), bold=False, spacing=0)
    if footer:
        _, tf = d._box(s, M, 6.5, CW, 0.4)
        d._run(d._para(tf, first=True), footer, 12, (THEME["mid_grey"] if dark else CORAL), bold=True, spacing=0.02)
    return s


def lifecycle(d, pre, accent, tag, stages, wrap, footer=None):
    """Three stage boxes (Implement › Optimise › Manage) with an engagement 'wrap' band
    underneath: a coral-outline segment spanning the first two stages + a coral-filled
    segment under the third. stages=[(name,line)x3]; wrap=((Lname,Ldesc),(Rname,Rdesc))."""
    s = d._slide(THEME["black"]); head(d, s, pre, accent, tag, dark=True)
    bw = (CW - 0.8) / 3; y = 2.4; h = 1.95
    for i, (name, line) in enumerate(stages[:3]):
        x = M + i * (bw + 0.4)
        d._rect(s, x, y, bw, h, fill_hex=DARK, line_hex=THEME["rev_hair"], line_w=1.0)
        _, tf = d._box(s, x + 0.3, y + 0.3, bw - 0.55, 0.4)
        d._run(d._para(tf, first=True), name.upper(), 14, CORAL, bold=True, spacing=0.08, caps=True)
        _, tf = d._box(s, x + 0.3, y + 0.84, bw - 0.6, h - 1.05)
        d._run(d._para(tf, first=True, line=1.32), line, 13.5, WHITE, bold=True, spacing=-0.01)
    for i in range(2):  # coral chevrons between the boxes
        ax = M + (i + 1) * (bw + 0.4) - 0.34
        _, tf = d._box(s, ax, y + h / 2 - 0.28, 0.6, 0.56, anchor=MSO_ANCHOR.MIDDLE)
        d._run(d._para(tf, first=True, align=PP_ALIGN.CENTER), "›", 26, CORAL, bold=True, spacing=0)
    (ll, ld), (rl, rd) = wrap
    wy = y + h + 0.45; wh = 1.0; lw = bw * 2 + 0.4
    d._rect(s, M, wy, lw, wh, fill_hex=None, line_hex=CORAL, line_w=1.5)
    _, tf = d._box(s, M + 0.3, wy + 0.2, lw - 0.6, 0.36)
    d._run(d._para(tf, first=True), ll.upper(), 11.5, CORAL, bold=True, spacing=0.1, caps=True)
    _, tf = d._box(s, M + 0.3, wy + 0.56, lw - 0.6, 0.36)
    d._run(d._para(tf, first=True), ld, 12, THEME["rev_body"], bold=False, spacing=0)
    rx = M + 2 * (bw + 0.4)
    d._rect(s, rx, wy, bw, wh, fill_hex=CORAL, line_hex=None)
    _, tf = d._box(s, rx + 0.3, wy + 0.2, bw - 0.6, 0.36)
    d._run(d._para(tf, first=True), rl.upper(), 11.5, WHITE, bold=True, spacing=0.1, caps=True)
    _, tf = d._box(s, rx + 0.3, wy + 0.56, bw - 0.6, 0.36)
    d._run(d._para(tf, first=True), rd, 12, WHITE, bold=True, spacing=0)
    if footer:
        _, tf = d._box(s, M, 6.72, CW, 0.4)
        d._run(d._para(tf, first=True), footer, 12, THEME["mid_grey"], bold=True, spacing=0.02)
    return s
