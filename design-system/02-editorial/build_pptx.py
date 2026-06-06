#!/usr/bin/env python3
# ============================================================================
# Mivada — Concept 02 · EDITORIAL  ·  python-pptx builder
#
# Bold magazine / thought-leadership. Black-dominant + coral, oversized
# Inter-800 titles, asymmetric layouts, big stat numerals, a coral pull-quote.
# Alternates black (#000) and off-white (#FAFAFA) slides for rhythm.
#
# A small THEME dict + a Deck class of helpers (cover, kpis, pillars, steps,
# split, pullquote, stats3, reasons, contact). A new deck is a few calls.
#
#   python build_pptx.py        ->  editorial.pptx
# ============================================================================

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.oxml.ns import qn
import os

# Real Mivada logos live in ../assets/logos (see its README for variants).
LOGO_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "logos")

# ---- THEME (exact Mivada tokens) -------------------------------------------
THEME = {
    "coral":      "EA493F",
    "coral_deep": "C9362B",
    "black":      "000000",
    "ink":        "111111",
    "off_white":  "FAFAFA",
    "white":      "FFFFFF",
    "warm":       "F2F3EE",
    "graphite":   "323232",
    "mid_grey":   "AEAEAE",
    "hairline":   "E4E4E0",
    # on-black body greys (tuned, never pure white)
    "rev_body":   "D6D6D6",
    "rev_soft":   "CFCFCF",
    "rev_hair":   "3A3A3A",
    "font":       "Inter",
}

EMU_IN = 914400
PAGE_W = 13.333
PAGE_H = 7.5

def C(name):
    return RGBColor.from_string(THEME[name])

def HEX(h):
    # accept either a 6-digit hex string or a THEME token name
    if isinstance(h, RGBColor):
        return h
    if h in THEME:
        return RGBColor.from_string(THEME[h])
    return RGBColor.from_string(h)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(PAGE_W)
        self.prs.slide_height = Inches(PAGE_H)
        self._blank = self.prs.slide_layouts[6]

    # ---- low-level helpers --------------------------------------------------
    def _slide(self, bg_hex):
        s = self.prs.slides.add_slide(self._blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = HEX(bg_hex)
        return s

    def _no_autofit(self, tf):
        # disable autosize so our point sizes are honoured exactly
        try:
            tf.word_wrap = True
            tf.auto_size = None
        except Exception:
            pass

    def _box(self, slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        self._no_autofit(tf)
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        return tb, tf

    def _run(self, para, text, size, color, bold=True, spacing=-0.02,
             font=None, italic=False, caps=False):
        r = para.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font or THEME["font"]
        r.font.color.rgb = color if isinstance(color, RGBColor) else HEX(color)
        r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_AUS
        # letter spacing (in points*100 attribute "spc")
        rPr = r._r.get_or_add_rPr()
        rPr.set("spc", str(int(spacing * size * 100)))
        if caps:
            rPr.set("cap", "all")
        return r

    def _para(self, tf, first=False, space_before=0, space_after=0,
              line=1.0, align=PP_ALIGN.LEFT):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        try:
            p.line_spacing = line
        except Exception:
            pass
        return p

    def _rect(self, slide, x, y, w, h, fill_hex=None, line_hex=None, line_w=None):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(x), Inches(y), Inches(w), Inches(h))
        shp.shadow.inherit = False
        if fill_hex is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = HEX(fill_hex)
        if line_hex is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = HEX(line_hex)
            shp.line.width = Pt(line_w or 1)
        return shp

    def _hline(self, slide, x, y, w, hex_, weight=1.0):
        ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = HEX(hex_)
        ln.line.width = Pt(weight)
        ln.shadow.inherit = False
        return ln

    def _vline(self, slide, x, y, h, hex_, weight=1.0):
        ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x), Inches(y + h))
        ln.line.color.rgb = HEX(hex_)
        ln.line.width = Pt(weight)
        ln.shadow.inherit = False
        return ln

    # ---- the "M" chip (coral rounded square + white Inter-800 M) -----------
    def _m_chip(self, slide, x, y, size_in):
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(x), Inches(y), Inches(size_in), Inches(size_in))
        chip.shadow.inherit = False
        chip.fill.solid()
        chip.fill.fore_color.rgb = C("coral")
        chip.line.fill.background()
        # ~20% corner radius
        try:
            chip.adjustments[0] = 0.20
        except Exception:
            pass
        tf = chip.text_frame
        self._no_autofit(tf)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._run(p, "M", size_in * 46, C("white"), bold=True, spacing=-0.04)
        return chip

    def _logo(self, slide, stem, x, y, height_in, fallback=True):
        """Place a real Mivada logo PNG (height-scaled, aspect preserved);
        fall back to the coral 'M' chip if the file is absent."""
        path = os.path.join(LOGO_DIR, stem + ".png")
        if os.path.exists(path):
            return slide.shapes.add_picture(path, Inches(x), Inches(y), height=Inches(height_in))
        if fallback:
            return self._m_chip(slide, x, y, height_in)
        return None

    def _eyebrow(self, slide, x, y, w, text, color="coral"):
        tb, tf = self._box(slide, x, y, w, 0.35)
        p = self._para(tf, first=True)
        self._run(p, text.upper(), 12, color, bold=True, spacing=0.16, caps=True)
        return tb

    def _slug(self, slide, text, dark=False):
        # small Mivada 'M' mark, top-left of every content slide
        self._logo(slide, "Mivada_Icon_White_RGB_L" if dark else "Mivada_Icon_Melon_RGB_L",
                   0.96, 0.5, 0.26, fallback=False)
        tb, tf = self._box(slide, PAGE_W - 4.3, 0.46, 3.5, 0.35)
        p = self._para(tf, first=True, align=PP_ALIGN.RIGHT)
        self._run(p, text.upper(), 11, "EA493F",
                  bold=True, spacing=0.14, caps=True)

    # ======================================================================
    # SLIDE HELPERS
    # ======================================================================

    def cover(self, brand, line_pre, line_accent, eyebrow, footL, footR):
        """(1) BLACK cover — oversized white headline, one coral word."""
        s = self._slide(THEME["black"])
        # real Mivada wordmark (coral M + white text, designed for black)
        self._logo(s, "Mivada_Logo_2C_OnBlack_RGB_L", 0.96, 0.82, 0.42)

        # eyebrow + monster headline
        self._eyebrow(s, 0.96, 2.55, 8, eyebrow, color="coral")
        tb, tf = self._box(s, 0.92, 2.95, 11.4, 3.4)
        p = self._para(tf, first=True, line=0.92)
        self._run(p, line_pre, 88, "FFFFFF", bold=True, spacing=-0.025)
        p2 = self._para(tf, line=0.92)
        self._run(p2, "human ", 88, "FFFFFF", bold=True, spacing=-0.025)
        self._run(p2, line_accent, 88, "EA493F", bold=True, spacing=-0.025)

        # foot rule + meta
        self._hline(s, 0.96, 6.74, PAGE_W - 1.92, THEME["rev_hair"], 1.0)
        tb, tf = self._box(s, 0.96, 6.86, 6, 0.4)
        p = self._para(tf, first=True)
        self._run(p, footL, 11.5, "AEAEAE", bold=True, spacing=0.04)
        tb, tf = self._box(s, PAGE_W - 7, 6.86, 6.04, 0.4)
        p = self._para(tf, first=True, align=PP_ALIGN.RIGHT)
        self._run(p, footR, 11.5, "AEAEAE", bold=True, spacing=0.04)
        return s

    def kpis(self, eyebrow, headline_pre, headline_accent, standfirst, stats):
        """(2) OFF-WHITE — oversized headline + standfirst, 3 big stats."""
        s = self._slide(THEME["off_white"])
        self._slug(s, "Who we are")

        # headline (left) + standfirst (right), baseline-ish aligned
        tb, tf = self._box(s, 0.92, 1.02, 7.7, 2.7)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, headline_pre + " ", 40, "111111", bold=True, spacing=-0.022)
        self._run(p, headline_accent, 40, "EA493F", bold=True, spacing=-0.022)

        tb, tf = self._box(s, 8.95, 2.05, 3.42, 2.4)
        p = self._para(tf, first=True, line=1.34)
        self._run(p, standfirst, 14.5, "323232", bold=False, spacing=0)

        # stat row along the bottom
        self._stat_row(s, 4.55, stats, light=True)
        return s

    def _stat_row(self, s, y, stats, light=True, top_rule=True):
        """3 big stat numerals across the width, vertical hairlines between."""
        n = len(stats)
        x0 = 0.96
        total = PAGE_W - 1.92
        col = total / n
        rule = THEME["hairline"] if light else THEME["rev_hair"]
        numc = "111111" if light else "FFFFFF"
        labc = THEME["mid_grey"]
        if top_rule:
            self._hline(s, x0, y - 0.18, total, rule, 1.0)
        for i, (num, accent, label) in enumerate(stats):
            cx = x0 + i * col
            if i > 0:
                self._vline(s, cx, y + 0.05, 1.7, rule, 1.0)
            pad = 0.0 if i == 0 else 0.34
            # adaptive numeral size: long word-values shrink to stay on one line
            chars = len(num) + (len(accent) if accent else 0)
            num_sz = 58 if chars <= 5 else (42 if chars <= 8 else 33)
            tb, tf = self._box(s, cx + pad, y, col - pad - 0.18, 1.15,
                               anchor=MSO_ANCHOR.BOTTOM)
            p = self._para(tf, first=True, line=0.92)
            self._run(p, num, num_sz, numc, bold=True, spacing=-0.035)
            if accent:
                self._run(p, accent, num_sz, "EA493F", bold=True, spacing=-0.035)
            tb, tf = self._box(s, cx + pad, y + 1.24, col - pad - 0.25, 0.7)
            p = self._para(tf, first=True, line=1.15)
            self._run(p, label.upper(), 10.5, labc, bold=True, spacing=0.12, caps=True)

    def pillars(self, eyebrow, headline_pre, headline_accent, items):
        """(3) BLACK — asymmetric editorial list, big coral index numerals."""
        s = self._slide(THEME["black"])
        self._slug(s, "What we do", dark=True)
        tb, tf = self._box(s, 0.92, 0.95, 11, 1.1)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, headline_pre + " ", 34, "FFFFFF", bold=True, spacing=-0.022)
        self._run(p, headline_accent, 34, "EA493F", bold=True, spacing=-0.022)

        y = 2.62
        rowh = 0.97
        x0 = 0.96
        w = PAGE_W - 1.92
        self._hline(s, x0, y, w, THEME["rev_hair"], 1.0)
        for i, (num, title, desc, tag) in enumerate(items):
            ry = y + i * rowh
            # index numeral
            tb, tf = self._box(s, x0, ry + 0.14, 1.1, 0.9)
            p = self._para(tf, first=True, line=1.0)
            self._run(p, num, 32, "EA493F", bold=True, spacing=-0.03)
            # title + desc
            tb, tf = self._box(s, x0 + 1.25, ry + 0.10, 7.7, 0.85)
            p = self._para(tf, first=True, line=1.0)
            self._run(p, title, 21, "FFFFFF", bold=True, spacing=-0.02)
            p2 = self._para(tf, space_before=4, line=1.12)
            self._run(p2, desc, 12, THEME["rev_soft"], bold=False, spacing=0)
            # tag (right)
            tb, tf = self._box(s, PAGE_W - 2.7, ry + 0.20, 1.74, 0.5, anchor=MSO_ANCHOR.TOP)
            p = self._para(tf, first=True, align=PP_ALIGN.RIGHT)
            self._run(p, tag.upper(), 10, THEME["mid_grey"], bold=True, spacing=0.13, caps=True)
            self._hline(s, x0, ry + rowh, w, THEME["rev_hair"], 1.0)
        return s

    def steps(self, eyebrow, headline_pre, headline_accent, steps):
        """(4) OFF-WHITE — bold numbered sequence, 4 across."""
        s = self._slide(THEME["off_white"])
        self._slug(s, "How we work")
        tb, tf = self._box(s, 0.92, 1.0, 11.4, 1.1)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, headline_pre + " ", 34, "111111", bold=True, spacing=-0.022)
        self._run(p, headline_accent, 34, "EA493F", bold=True, spacing=-0.022)

        y = 3.25
        x0 = 0.96
        total = PAGE_W - 1.92
        col = total / len(steps)
        for i, (n, name, desc) in enumerate(steps):
            cx = x0 + i * col
            if i > 0:
                self._vline(s, cx, y, 2.5, THEME["hairline"], 1.0)
            pad = 0.0 if i == 0 else 0.34
            tb, tf = self._box(s, cx + pad, y, col - pad - 0.25, 2.5)
            p = self._para(tf, first=True)
            self._run(p, n.upper(), 11, "EA493F", bold=True, spacing=0.16, caps=True)
            p2 = self._para(tf, space_before=10, line=1.0)
            self._run(p2, name, 30, "111111", bold=True, spacing=-0.025)
            p3 = self._para(tf, space_before=10, line=1.28)
            self._run(p3, desc, 12.5, "323232", bold=False, spacing=0)
        return s

    def split(self, eyebrow, headline_pre, headline_accent, standfirst, caps):
        """(5) BLACK — Data & AI: headline+standfirst left, capability list right."""
        s = self._slide(THEME["black"])
        self._slug(s, "Data & AI", dark=True)
        # left column
        tb, tf = self._box(s, 0.92, 1.2, 5.4, 1.9)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, headline_pre + " ", 36, "FFFFFF", bold=True, spacing=-0.022)
        self._run(p, headline_accent, 36, "EA493F", bold=True, spacing=-0.022)
        tb, tf = self._box(s, 0.96, 3.95, 5.0, 2.2)
        p = self._para(tf, first=True, line=1.34)
        self._run(p, standfirst, 14, THEME["rev_soft"], bold=False, spacing=0)

        # right column — capability list (key over value, full width = no overlap)
        rx = 6.95
        rw = PAGE_W - 0.96 - rx
        y = 1.35
        rowh = 1.2
        self._hline(s, rx, y, rw, THEME["rev_hair"], 1.0)
        for i, (k, v_parts) in enumerate(caps):
            ry = y + i * rowh
            tb, tf = self._box(s, rx, ry + 0.2, rw, 0.32)
            p = self._para(tf, first=True)
            self._run(p, k.upper(), 10.5, "EA493F", bold=True, spacing=0.14, caps=True)
            tb, tf = self._box(s, rx, ry + 0.54, rw, 0.5)
            p = self._para(tf, first=True, line=1.1)
            for txt, strong in v_parts:
                self._run(p, txt, 14, "FFFFFF" if strong else THEME["rev_body"],
                          bold=strong, spacing=-0.005)
            self._hline(s, rx, ry + rowh, rw, THEME["rev_hair"], 1.0)
        return s

    def pullquote(self, eyebrow, quote, attr, stats):
        """(6) OFF-WHITE — big coral pull-quote + 3 big stats."""
        s = self._slide(THEME["off_white"])
        self._slug(s, "Outcomes")
        # coral left rule for the quote
        self._rect(s, 0.96, 1.55, 0.07, 2.5, fill_hex=THEME["coral"])
        tb, tf = self._box(s, 1.32, 1.5, 10.5, 2.7)
        p = self._para(tf, first=True, line=1.06)
        self._run(p, quote, 38, "EA493F", bold=True, spacing=-0.02)
        tb, tf = self._box(s, 1.32, 4.18, 10, 0.4)
        p = self._para(tf, first=True)
        self._run(p, attr, 12, THEME["mid_grey"], bold=True, spacing=0.02)

        self._stat_row(s, 5.35, stats, light=True)
        return s

    def reasons(self, eyebrow, headline_pre, headline_accent, items):
        """(7) BLACK — 3 bold reasons across."""
        s = self._slide(THEME["black"])
        self._slug(s, "Why Mivada", dark=True)
        tb, tf = self._box(s, 0.92, 1.0, 11, 1.1)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, headline_pre + " ", 34, "FFFFFF", bold=True, spacing=-0.022)
        self._run(p, headline_accent, 34, "EA493F", bold=True, spacing=-0.022)

        y = 3.25
        x0 = 0.96
        total = PAGE_W - 1.92
        col = total / len(items)
        for i, (n, h, b) in enumerate(items):
            cx = x0 + i * col
            if i > 0:
                self._vline(s, cx, y, 2.55, THEME["rev_hair"], 1.0)
            pad = 0.0 if i == 0 else 0.36
            tb, tf = self._box(s, cx + pad, y, col - pad - 0.3, 2.6)
            p = self._para(tf, first=True)
            self._run(p, n, 13, "EA493F", bold=True, spacing=0.14)
            p2 = self._para(tf, space_before=12, line=1.02)
            self._run(p2, h, 23, "FFFFFF", bold=True, spacing=-0.025)
            p3 = self._para(tf, space_before=12, line=1.28)
            self._run(p3, b, 12.5, THEME["rev_soft"], bold=False, spacing=0)
        return s

    def contact(self, brand, eyebrow, line_pre, line_accent, footL, footR):
        """(8) BLACK closing — oversized coral/white CTA + M-chip."""
        s = self._slide(THEME["black"])
        self._logo(s, "Mivada_Logo_2C_OnBlack_RGB_L", 0.96, 0.82, 0.42)

        self._eyebrow(s, 0.96, 2.4, 8, eyebrow, color="coral")
        tb, tf = self._box(s, 0.92, 2.85, 11.0, 3.0)
        p = self._para(tf, first=True, line=0.98)
        self._run(p, line_pre + " ", 54, "FFFFFF", bold=True, spacing=-0.025)
        self._run(p, line_accent, 54, "EA493F", bold=True, spacing=-0.025)

        self._hline(s, 0.96, 6.62, PAGE_W - 1.92, THEME["rev_hair"], 1.0)
        tb, tf = self._box(s, 0.96, 6.76, 7, 0.5)
        p = self._para(tf, first=True)
        self._run(p, footL, 12.5, "CFCFCF", bold=True, spacing=0.02)
        tb, tf = self._box(s, PAGE_W - 7, 6.76, 6.04, 0.5)
        p = self._para(tf, first=True, align=PP_ALIGN.RIGHT)
        self._run(p, footR, 12.5, "AEAEAE", bold=False, spacing=0.02)
        return s

    # ======================================================================
    # GENERAL-PURPOSE slides (for translating arbitrary decks — never lose text)
    # ======================================================================
    def _fit(self, text, big, mid, small, a=18, b=30):
        n = len(text or "")
        return big if n <= a else (mid if n <= b else small)

    def cover_plain(self, eyebrow, headline_pre, headline_accent, footL="", footR="",
                    logo="Mivada_Logo_2C_OnBlack_RGB_L"):
        """BLACK cover with a free headline (pre + coral accent)."""
        s = self._slide(THEME["black"])
        self._logo(s, logo, 0.96, 0.82, 0.42)
        if eyebrow:
            self._eyebrow(s, 0.96, 2.55, 10.5, eyebrow, color="coral")
        sz = self._fit((headline_pre or "") + (headline_accent or ""), 76, 54, 40)
        _, tf = self._box(s, 0.92, 2.95, 11.6, 3.2)
        p = self._para(tf, first=True, line=0.98)
        self._run(p, (headline_pre or "") + (" " if headline_accent else ""), sz, "FFFFFF", bold=True, spacing=-0.025)
        if headline_accent:
            self._run(p, headline_accent, sz, "EA493F", bold=True, spacing=-0.025)
        self._hline(s, 0.96, 6.74, PAGE_W - 1.92, THEME["rev_hair"], 1.0)
        _, tf = self._box(s, 0.96, 6.86, 7.5, 0.4)
        self._run(self._para(tf, first=True), footL or "", 11.5, "AEAEAE", bold=True, spacing=0.04)
        _, tf = self._box(s, PAGE_W - 7, 6.86, 6.04, 0.4)
        self._run(self._para(tf, first=True, align=PP_ALIGN.RIGHT), footR or "", 11.5, "AEAEAE", bold=True, spacing=0.04)
        return s

    def contact_plain(self, eyebrow, headline_pre, headline_accent, footL="", footR="",
                      logo="Mivada_Logo_2C_OnBlack_RGB_L"):
        """BLACK closing with a free headline + contact feet."""
        s = self._slide(THEME["black"])
        self._logo(s, logo, 0.96, 0.82, 0.42)
        if eyebrow:
            self._eyebrow(s, 0.96, 2.4, 10.5, eyebrow, color="coral")
        sz = self._fit((headline_pre or "") + (headline_accent or ""), 54, 44, 34, a=26, b=46)
        _, tf = self._box(s, 0.92, 2.85, 11.2, 3.0)
        p = self._para(tf, first=True, line=0.98)
        self._run(p, (headline_pre or "") + (" " if headline_accent else ""), sz, "FFFFFF", bold=True, spacing=-0.025)
        if headline_accent:
            self._run(p, headline_accent, sz, "EA493F", bold=True, spacing=-0.025)
        self._hline(s, 0.96, 6.4, PAGE_W - 1.92, THEME["rev_hair"], 1.0)
        _, tf = self._box(s, 0.96, 6.55, 8, 0.5)
        self._run(self._para(tf, first=True), footL or "", 12.5, "CFCFCF", bold=True, spacing=0.02)
        _, tf = self._box(s, PAGE_W - 7, 6.55, 6.04, 0.5)
        self._run(self._para(tf, first=True, align=PP_ALIGN.RIGHT), footR or "", 12.5, "AEAEAE", bold=False, spacing=0.02)
        return s

    def content(self, eyebrow, headline_pre, headline_accent, items, dark=False, slug_txt=None):
        """A general content slide — title + a left-aligned coral-bulleted body.
        Holds an arbitrary number of lines (auto-shrinks) so no source text is dropped."""
        s = self._slide(THEME["black"] if dark else THEME["off_white"])
        self._slug(s, (slug_txt or eyebrow or "")[:34], dark=dark)
        if eyebrow:
            self._eyebrow(s, 0.96, 0.92, 11, eyebrow, color="coral")
        tcol = "FFFFFF" if dark else "111111"
        hz = self._fit((headline_pre or "") + (headline_accent or ""), 34, 28, 23, a=26, b=46)
        _, tf = self._box(s, 0.92, 1.34, 11.4, 1.2)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, (headline_pre or "") + (" " if headline_accent else ""), hz, tcol, bold=True, spacing=-0.022)
        if headline_accent:
            self._run(p, headline_accent, hz, "EA493F", bold=True, spacing=-0.022)
        items = [it for it in (items or []) if it and str(it).strip()]
        n = len(items)
        bsize = 16 if n <= 5 else (13 if n <= 9 else 11)
        gap = 9 if n <= 5 else 5
        body = THEME["rev_body"] if dark else "323232"
        _, tf = self._box(s, 0.96, 2.75, PAGE_W - 1.92, 4.35)
        for i, it in enumerate(items):
            p = self._para(tf, first=(i == 0), space_before=(0 if i == 0 else gap), line=1.22)
            self._run(p, "—  ", bsize, "EA493F", bold=True, spacing=0)
            self._run(p, str(it), bsize, body, bold=False, spacing=0)
        return s

    def quote_plain(self, eyebrow, quote, attr="", dark=False):
        """A general pull-quote slide (no stats required)."""
        s = self._slide(THEME["black"] if dark else THEME["off_white"])
        self._slug(s, (eyebrow or "Quote")[:34], dark=dark)
        if eyebrow:
            self._eyebrow(s, 0.96, 0.92, 10, eyebrow, color="coral")
        self._rect(s, 0.96, 1.7, 0.07, 3.2, fill_hex=THEME["coral"])
        qs = 36 if len(quote or "") < 140 else (28 if len(quote or "") < 240 else 21)
        _, tf = self._box(s, 1.32, 1.6, 10.8, 3.8)
        self._run(self._para(tf, first=True, line=1.1), quote or "", qs, "EA493F", bold=True, spacing=-0.02)
        if attr:
            _, tf = self._box(s, 1.32, 5.5, 10, 0.5)
            self._run(self._para(tf, first=True), attr, 12, THEME["mid_grey"], bold=True, spacing=0.02)
        return s

    def chart_coral(self, eyebrow, headline_pre, headline_accent, categories, series, slug_txt=None):
        """A column chart re-coloured to the brand (coral primary), data preserved."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        s = self._slide(THEME["off_white"])
        self._slug(s, (slug_txt or eyebrow or "Results")[:34])
        if eyebrow:
            self._eyebrow(s, 0.96, 0.92, 9, eyebrow, color="coral")
        _, tf = self._box(s, 0.92, 1.34, 11, 1.0)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, (headline_pre or "") + (" " if headline_accent else ""), 32, "111111", bold=True, spacing=-0.022)
        if headline_accent:
            self._run(p, headline_accent, 32, "EA493F", bold=True, spacing=-0.022)
        cd = CategoryChartData(); cd.categories = categories
        for name, vals in series:
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                Inches(0.96), Inches(2.7), Inches(11.4), Inches(4.0), cd)
        ch = gf.chart; ch.has_legend = len(series) > 1
        plot = ch.plots[0]; plot.gap_width = 70
        palette = ["EA493F", "111111", "AEAEAE", "C9362B"]
        for i, se in enumerate(plot.series):
            se.format.fill.solid(); se.format.fill.fore_color.rgb = RGBColor.from_string(palette[i % len(palette)])
            se.format.line.fill.background()
        if len(series) == 1:
            plot.has_data_labels = True
            plot.data_labels.number_format = "0"; plot.data_labels.number_format_is_linked = False
        return s

    def media(self, eyebrow, headline_pre, headline_accent, items, image_path,
              dark=False, slug_txt=None):
        """Split slide — text/bullets left, an image placed (aspect-preserved) right.
        Used so images from a translated deck are carried, not dropped."""
        s = self._slide(THEME["black"] if dark else THEME["off_white"])
        self._slug(s, (slug_txt or eyebrow or "")[:30], dark=dark)
        if eyebrow:
            self._eyebrow(s, 0.96, 0.92, 6.0, eyebrow, color="coral")
        tcol = "FFFFFF" if dark else "111111"
        _, tf = self._box(s, 0.92, 1.34, 5.9, 1.3)
        p = self._para(tf, first=True, line=1.0)
        self._run(p, (headline_pre or "") + (" " if headline_accent else ""), 30, tcol, bold=True, spacing=-0.022)
        if headline_accent:
            self._run(p, headline_accent, 30, "EA493F", bold=True, spacing=-0.022)
        items = [it for it in (items or []) if it and str(it).strip()]
        bsize = 14 if len(items) <= 6 else 11
        body = THEME["rev_body"] if dark else "323232"
        _, tf = self._box(s, 0.96, 2.75, 5.5, 4.3)
        for i, it in enumerate(items):
            p = self._para(tf, first=(i == 0), space_before=(0 if i == 0 else 7), line=1.22)
            self._run(p, "—  ", bsize, "EA493F", bold=True, spacing=0)
            self._run(p, str(it), bsize, body, bold=False, spacing=0)
        if image_path and os.path.exists(image_path):
            bx, by, bw, bh = 6.95, 1.45, 5.42, 5.0
            try:
                from PIL import Image as _PIL
                iw, ih = _PIL.open(image_path).size; ar = iw / ih
            except Exception:
                ar = 1.5
            w = bw; h = w / ar
            if h > bh:
                h = bh; w = h * ar
            s.shapes.add_picture(image_path, Inches(bx + (bw - w) / 2),
                                 Inches(by + (bh - h) / 2), width=Inches(w))
        return s

    def save(self, path):
        self.prs.save(path)


# ============================================================================
# Build the 8-slide deck (content only — the look lives in the helpers above)
# ============================================================================
def build():
    d = Deck()

    # (1) Cover — BLACK
    d.cover(
        "Mivada",
        "Technology,", "first.",
        "Capability overview · 2026",
        "An Australian technology consultancy",
        "Workday · Payroll · Data & AI · Automation",
    )

    # (2) Who we are — OFF-WHITE
    d.kpis(
        "Who we are",
        "We turn enterprise platforms into measurable",
        "human outcomes.",
        "Mivada takes Workday, payroll, data and AI and shapes them around how your "
        "people actually work — onshore in Australia, with delivery across India.",
        [
            ("2014", "", "Founded · formerly LJM Infotech"),
            ("120", "+", "Specialists, not generalists"),
            ("AU + India", "", "Onshore lead · offshore delivery"),
        ],
    )

    # (3) What we do — BLACK
    d.pillars(
        "Four service pillars",
        "Four practices, one", "operating belief.",
        [
            ("01", "Workday & ERP",
             "Implementation, optimisation and managed support for Workday HCM, Financials and adjacent ERP.",
             "Platform"),
            ("02", "Payroll consulting",
             "Compliant, accurate payroll across complex awards and multi-entity structures.",
             "Compliance"),
            ("03", "Data & AI",
             "Lakehouse architecture, governed analytics and applied AI on your people and finance data.",
             "Insight"),
            ("04", "Intelligent automation",
             "RPA, ML and process design combined to remove low-value work.",
             "Efficiency"),
        ],
    )

    # (4) How we work — OFF-WHITE
    d.steps(
        "Think human first",
        "Understand the people", "before the technology.",
        [
            ("Step 01", "Listen", "Understand the people and the work before a line of configuration."),
            ("Step 02", "Design", "Shape the platform around how teams actually operate."),
            ("Step 03", "Deliver", "Implement in weeks-long increments, not multi-year programs."),
            ("Step 04", "Care", "Stay on after go-live; measure outcomes, not tickets closed."),
        ],
    )

    # (5) Data & AI — BLACK
    d.split(
        "Capability spotlight",
        "Decisions made on", "current data.",
        "Certified consultants build the data foundation that puts a governed, real-time "
        "view of your people and finance data in front of leaders.",
        [
            ("Architecture", [("Designed ", False), ("Lakehouse architectures", True),
                              (" that unify people and finance data.", False)]),
            ("Pipelines", [("Built ", False), ("Delta Lake pipelines", True),
                           (" — batch and streaming — on ", False), ("Databricks", True), (".", False)]),
            ("Integration", [("Connected Databricks to ", False),
                             ("Workday, payroll and finance", True), (".", False)]),
            ("Outcome", [("Delivered ", False), ("governed KPI stores", True),
                         (" and ", False), ("real-time insight", True), (".", False)]),
        ],
    )

    # (6) Outcomes — OFF-WHITE pull-quote + stats
    d.pullquote(
        "In their words",
        "“They delivered in weeks what we'd scoped for a year — and our team actually uses it.”",
        "People & Payroll Director · national employer (illustrative)",
        [
            ("40", "+", "Workday deployments"),
            ("98", "%", "Client retention"),
            ("Weeks", "", "To value · not months"),
        ],
    )

    # (7) Why Mivada — BLACK
    d.reasons(
        "Three reasons",
        "Why teams choose to", "stay with us.",
        [
            ("01", "Certified specialists",
             "Workday- and Databricks-certified consultants, not generalists learning on your budget."),
            ("02", "Onshore + offshore",
             "Senior onshore leadership in Australia with cost-effective delivery from India."),
            ("03", "People-first change",
             "Adoption built in from day one, so the platform actually gets used."),
        ],
    )

    # (8) Contact — BLACK closing
    d.contact(
        "Mivada",
        "Let's talk",
        "Let's build smarter systems, with people at the", "centre.",
        "hello@mivada.com · mivada.com",
        "Sydney & Melbourne · Onshore AU & India",
    )

    d.save("editorial.pptx")
    print("Wrote editorial.pptx (%d slides)" % len(d.prs.slides._sldIdLst))


if __name__ == "__main__":
    build()
