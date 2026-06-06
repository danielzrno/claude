#!/usr/bin/env python3
# ============================================================================
# Mivada — Company Overview + AMS & Services, merged into one Editorial deck.
# Overview slides + the shared AMS/services builders (ams_slides.py), which keep
# the original diagram structures (framework stack, gantt, pyramid, org chart).
# Single coral section label per slide. 17 slides.
#   python mivada_full_deck.py -> Mivada_Overview_and_Services.pptx
# ============================================================================
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "02-editorial"))
from build_pptx import Deck, THEME, PAGE_W
from ams_slides import (head, standfirst, fin_hcm, engagement_models,
                        ams_framework, coverage_model, governance_model, team_structure)
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from PIL import Image as PILImage

CORAL, INK, WHITE = "EA493F", "111111", "FFFFFF"
M = 0.96; CW = PAGE_W - 2 * M
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "logos", "Mivada_Logo_Master_RGB_L.png")
d = Deck()

# 1 · LOGO (white) -----------------------------------------------------------
s = d._slide(WHITE)
_h = 0.98; _w = _h * (1002 / 152)
s.shapes.add_picture(LOGO, Inches((PAGE_W - _w) / 2), Inches(2.72), height=Inches(_h))
_, tf = d._box(s, 2, 4.12, PAGE_W - 4, 0.6); p = d._para(tf, first=True, align=PP_ALIGN.CENTER)
d._run(p, "Technology, ", 16, INK, bold=True, spacing=0); d._run(p, "human first.", 16, CORAL, bold=True, spacing=0)

# 2 · COVER ------------------------------------------------------------------
d.cover_plain("Mivada · Company overview · 2026", "Technology,", "human first.",
              "Australian technology consultancy", "mivada.com")

# 3 · WHO WE ARE -------------------------------------------------------------
d.kpis("Who we are", "An Australian technology", "consultancy.",
       "We turn Workday, people systems, data and automation into outcomes teams actually feel — "
       "where human understanding becomes system advantage.",
       [("2014", "", "Founded"), ("200", "+", "Specialists · AU, US & India"),
        ("100", "+", "People-systems clients")])

# 4 · TRUSTED BY -------------------------------------------------------------
def clients_slide():
    img = next((c for c in ["../logos.png", "../assets/clients/clients-logos.png",
                            "../assets/clients/clients.png"] if os.path.exists(c)), None)
    if img:
        s = d._slide(WHITE); d._slug(s, "Trusted by")
        _, tf = d._box(s, 0.92, 1.0, CW, 1.0); p = d._para(tf, first=True)
        d._run(p, "In good ", 33, INK, bold=True, spacing=-0.022); d._run(p, "company.", 33, CORAL, bold=True, spacing=-0.022)
        iw, ih = PILImage.open(img).size; w = 10.8; h = w * ih / iw
        if h > 3.9: h = 3.9; w = h * iw / ih
        s.shapes.add_picture(img, Inches((PAGE_W - w) / 2), Inches(2.7), width=Inches(w)); return
    s = d._slide(THEME["off_white"]); head(d, s, "In good", "company.", "Trusted by")
    standfirst(d, s, "A decade of work with some of Australia's most demanding operations.", 2.05)
    names = ["Qantas", "Jetstar", "Guzman y Gomez", "Western Sydney Airport", "Canva", "NAB",
             "Rio Tinto", "ResMed", "HCF", "IAG", "University of Sydney", "Macquarie University",
             "Nine", "Seven West Media", "Amart", "dnata", "Queensland Airports", "Kennards",
             "Anglicare", "ProPharma"]
    cols = 4; cw = CW / cols; y0 = 2.95; rowh = 0.62
    for i, nm in enumerate(names):
        r, c = divmod(i, cols); x = M + c * cw; y = y0 + r * rowh
        _, tf = d._box(s, x, y, cw - 0.2, 0.5); d._run(d._para(tf, first=True), nm, 15, INK, bold=True, spacing=-0.01)
        d._hline(s, x, y + 0.5, cw - 0.35, THEME["hairline"], 1.0)
clients_slide()

# 5 · WHAT WE DO -------------------------------------------------------------
d.pillars("What we do", "Across your", "people platform.", [
    ("01", "People Systems & Transformation", "HCM, Payroll, Time & Attendance and HRIS — implemented and optimised.", "Workday"),
    ("02", "Data, Analytics & BI", "Reporting, dashboards and governed analytics leaders trust.", "Insight"),
    ("03", "Process & Automation", "Redesign, integrations and automation that remove manual work.", "Efficiency"),
    ("04", "Managed Services & Support", "Long-term AMS that keeps delivering value after go-live.", "Partnership"),
])

# 6 · HOW WE WORK ------------------------------------------------------------
d.steps("How we work", "People first,", "start to finish.", [
    ("01", "Listen", "We learn your people, processes and constraints before touching the platform."),
    ("02", "Design", "We shape Workday around how your teams really work."),
    ("03", "Deliver", "Agile delivery, senior practitioners, clear governance."),
    ("04", "Care", "Long-term managed services that keep improving outcomes."),
])

# 7 · DATA & AI --------------------------------------------------------------
d.split("Data & AI", "Decisions on", "current data.",
        "From reporting to governed analytics and automation — we turn Workday data into trusted, timely insight.",
        [("Report", [("Dashboards & reporting people actually use", True)]),
         ("Analyse", [("Prism & data lakehouse, governed analytics", True)]),
         ("Automate", [("Integrations & automation across the stack", True)]),
         ("Assure", [("Security, controls and audit readiness", True)])])

# 8-13 · SERVICES DEPTH (shared builders — keep original diagram structures) --
fin_hcm(d)
engagement_models(d)
ams_framework(d)
coverage_model(d)
governance_model(d)
team_structure(d)

# 14 · OUTCOMES (case studies) ----------------------------------------------
s = d._slide(THEME["off_white"]); head(d, s, "Proven Workday", "outcomes.", "Outcomes")
standfirst(d, s, "Long-term partnerships with some of Australia's most demanding operations.", 2.0, size=13)
cases = [
    ("Qantas", "A partner since 2014 — taking over Workday support in 2017 and growing into a strategic partner across HR, Payroll, Digital, Automation, Data and Analytics. We scaled integrations, modernised architectures and embedded automation, evolving into a long-term enterprise partner."),
    ("Guzman y Gomez", "Stabilised and optimised Workday post–go-live across modules, reporting, security and compliance. Agile delivery onboarded 150+ franchisees, shipped new reporting and a redesigned security model, and lifted performance across HCM, Absence and Learning."),
    ("Western Sydney Intl Airport", "Delivered a card-management system, then expanded into a strategic role on WSI's Workday program. By resolving functional and integration issues and assuming full AMS ownership, we stabilised HR, Payroll, Finance and Procurement.")]
col = CW / 3
for i, (client, para) in enumerate(cases):
    cx = M + i * col
    if i > 0: d._vline(s, cx, 2.95, 3.8, THEME["hairline"], 1.0)
    pad = 0.0 if i == 0 else 0.32
    _, tf = d._box(s, cx + pad, 2.9, col - pad - 0.25, 0.5); d._run(d._para(tf, first=True), client.upper(), 12, CORAL, bold=True, spacing=0.06, caps=True)
    _, tf = d._box(s, cx + pad, 3.45, col - pad - 0.28, 3.4); d._run(d._para(tf, first=True, line=1.28), para, 10.5, "323232", bold=False, spacing=0)

# 15 · WHY MIVADA ------------------------------------------------------------
d.reasons("Why Mivada", "A partner,", "not a vendor.", [
    ("01", "We grew out of a customer", "So we understand your challenges from the inside — and we're not a Big 4."),
    ("02", "Senior, local, accountable", "Australian-owned, with direct access to our CEO and a dedicated engagement manager."),
    ("03", "In it for the long term", "From go-live into managed services — we grow with you.")])

# 16 · CAPABILITY ------------------------------------------------------------
s = d._slide(THEME["off_white"]); head(d, s, "Proven in", "Workday.", "Capability")
standfirst(d, s, "A decade of Workday delivery — consultants average 3+ years on the platform.", 2.05)
STATS = [("3", "+", "Certifications / consultant"), ("2000", "+", "Workday integrations"),
         ("3000", "+", "Custom reports"), ("30", "+", "Implementations & support"),
         ("200", "+", "Employees · AU, US & India"), ("100", "+", "People-systems clients"),
         ("10", "+", "Years delivering Workday")]
d._stat_row(s, 3.6, STATS[:4], light=True)
d._stat_row(s, 5.6, STATS[4:], light=True)

# 17 · CONTACT ---------------------------------------------------------------
d.contact_plain("Let's talk", "Let's build something", "human.",
                "hello@mivada.com · mivada.com", "Sydney · Melbourne · AU & India")

OUT = "Mivada_Overview_and_Services.pptx"
d.save(OUT)
print(f"Wrote {OUT} ({len(d.prs.slides._sldIdLst)} slides)")
