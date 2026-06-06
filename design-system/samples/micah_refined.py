#!/usr/bin/env python3
# ============================================================================
# First meeting with Micah Projects — REFINED single-version Editorial deck.
# Collapses the triplicated sections to the strongest version of each, removes
# repeated messaging (pillars -> Services; numbers -> Capability; trusted-by
# once), reorders for a compelling first-meeting narrative, adds a white logo
# slide + a reworked intro + a "how we deliver" divider + a next-steps close.
#   python micah_refined.py -> First_Meeting_with_Micah_Projects_refined.pptx
# ============================================================================
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "02-editorial"))
from build_pptx import Deck, THEME, PAGE_W, LOGO_DIR
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image as PILImage

CORAL, INK, WHITE = "EA493F", "111111", "FFFFFF"
M = 0.96; CW = PAGE_W - 2 * M
BRANDS = "Qantas   ·   Guzman y Gomez   ·   Western Sydney International Airport"
src = Presentation("First_Meeting_with_Micah_Projects.pptx")
d = Deck()

# ---------- helpers ---------------------------------------------------------
def head(s, eyebrow, hpre, hacc, tag, dark=False, hsize=34, y=1.32):
    d._slug(s, tag, dark=dark)
    d._eyebrow(s, M, 0.92, 11, eyebrow, color="coral")
    _, tf = d._box(s, 0.92, y, CW, 1.2); p = d._para(tf, first=True, line=1.0)
    d._run(p, hpre + (" " if hacc else ""), hsize, (WHITE if dark else INK), bold=True, spacing=-0.022)
    if hacc: d._run(p, hacc, hsize, CORAL, bold=True, spacing=-0.022)

def standfirst(s, text, y, w=10.8, color="323232", size=14):
    _, tf = d._box(s, M, y, w, 0.9)
    d._run(d._para(tf, first=True, line=1.3), text, size, color, bold=False, spacing=0)

def trusted(s, y=6.92, dark=False):
    d._hline(s, M, y - 0.16, CW, THEME["rev_hair"] if dark else THEME["hairline"], 1.0)
    _, tf = d._box(s, M, y, CW, 0.4); p = d._para(tf, first=True)
    d._run(p, "TRUSTED BY   ", 10.5, CORAL, bold=True, spacing=0.14, caps=True)
    d._run(p, BRANDS, 10.5, THEME["rev_soft"] if dark else THEME["mid_grey"], bold=True, spacing=0.04)

def card(s, x, y, w, h, num, title, desc=None, dark=False):
    d._rect(s, x, y, w, h, fill_hex=("141414" if dark else WHITE),
            line_hex=(THEME["rev_hair"] if dark else THEME["hairline"]), line_w=1.0)
    _, tf = d._box(s, x + 0.26, y + 0.2, w - 0.45, 0.4)
    d._run(d._para(tf, first=True), num, 13.5, CORAL, bold=True, spacing=0.06)
    _, tf = d._box(s, x + 0.26, y + 0.54, w - 0.45, 0.62)
    d._run(d._para(tf, first=True, line=1.04), title, 14.5, (WHITE if dark else INK), bold=True, spacing=-0.01)
    if desc:
        _, tf = d._box(s, x + 0.26, y + 1.14, w - 0.48, h - 1.2)
        d._run(d._para(tf, first=True, line=1.16), desc, 10.5,
               (THEME["rev_soft"] if dark else "323232"), bold=False, spacing=0)

def grid(s, items, cols, x0, y0, w, rowh, gapx, gapy, dark=False, desc=True):
    cw = (w - gapx * (cols - 1)) / cols
    for i, it in enumerate(items):
        r, c = divmod(i, cols); x = x0 + c * (cw + gapx); y = y0 + r * (rowh + gapy)
        card(s, x, y, cw, rowh, it[0], it[1], (it[2] if desc and len(it) > 2 else None), dark=dark)

def imgs_from(idx, n):
    pics = sorted([sh for sh in src.slides[idx].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE],
                  key=lambda s: s.width * s.height, reverse=True)
    out = []
    for k, sh in enumerate(pics[:n]):
        ct = sh.image.content_type; ext = ".wmf" if "wmf" in ct else (".jpg" if "jpeg" in ct else ".png")
        raw = f"/tmp/mr_{idx}_{k}{ext}"; open(raw, "wb").write(sh.image.blob); out.append(raw)
    return out

def place_images(s, paths, x=0.96, y=2.25, w=11.4, h=4.6):
    n = len(paths)
    if n == 0: return
    gap = 0.3 if n > 1 else 0; cw = (w - gap * (n - 1)) / n; cx = x
    for p in paths:
        try: iw, ih = PILImage.open(p).size; ar = iw / ih
        except Exception: ar = 1.4
        ww = cw; hh = ww / ar
        if hh > h: hh = h; ww = hh * ar
        try: s.shapes.add_picture(p, Inches(cx + (cw - ww) / 2), Inches(y + (h - hh) / 2), width=Inches(ww))
        except Exception as e: print("  ! img skipped", e)
        cx += cw + gap

def diagram(eyebrow, hpre, hacc, idx, n=1, tag=None):
    s = d._slide(THEME["off_white"]); head(s, eyebrow, hpre, hacc, tag or eyebrow)
    place_images(s, imgs_from(idx, n)); return s

def divider(eyebrow, tpre, tacc, sub):
    s = d._slide(THEME["black"]); d._slug(s, eyebrow, dark=True)
    d._eyebrow(s, M, 2.5, 8, eyebrow, color="coral")
    _, tf = d._box(s, 0.92, 2.95, 11.4, 1.6); p = d._para(tf, first=True, line=0.98)
    d._run(p, tpre + " ", 52, WHITE, bold=True, spacing=-0.025); d._run(p, tacc, 52, CORAL, bold=True, spacing=-0.025)
    if sub:
        _, tf = d._box(s, 0.96, 4.5, 10.5, 0.6)
        d._run(d._para(tf, first=True), sub, 15, THEME["rev_body"], bold=False, spacing=0)
    return s

# ============================================================================
# 1 · LOGO — white background
# ============================================================================
s = d._slide(WHITE)
_h = 0.98; _ar = 1002 / 152; _w = _h * _ar
s.shapes.add_picture(os.path.join(LOGO_DIR, "Mivada_Logo_Master_RGB_L.png"),
                     Inches((PAGE_W - _w) / 2), Inches(2.72), height=Inches(_h))
_, tf = d._box(s, 2, 4.12, PAGE_W - 4, 0.6); p = d._para(tf, first=True, align=PP_ALIGN.CENTER)
d._run(p, "Technology, ", 16, INK, bold=True, spacing=0.0); d._run(p, "human first.", 16, CORAL, bold=True, spacing=0.0)

# ============================================================================
# 2 · COVER — framed for the client (black)
# ============================================================================
s = d._slide(THEME["black"])
d._logo(s, "Mivada_Logo_2C_OnBlack_RGB_L", 0.96, 0.82, 0.42)
d._eyebrow(s, 0.96, 2.5, 10, "Micah Projects · First meeting · March 2026", color="coral")
_, tf = d._box(s, 0.92, 2.95, 11.6, 2.6); p = d._para(tf, first=True, line=0.94)
d._run(p, "A partner for your", 70, WHITE, bold=True, spacing=-0.028)
p2 = d._para(tf, line=0.94); d._run(p2, "Workday ", 70, WHITE, bold=True, spacing=-0.028); d._run(p2, "journey.", 70, CORAL, bold=True, spacing=-0.028)
d._hline(s, 0.96, 6.74, CW, THEME["rev_hair"], 1.0)
_, tf = d._box(s, 0.96, 6.86, 7, 0.4); d._run(d._para(tf, first=True), "Mivada — Workday implementation & managed services", 11.5, "AEAEAE", bold=True, spacing=0.04)
_, tf = d._box(s, PAGE_W - 7, 6.86, 6.04, 0.4); p = d._para(tf, first=True, align=PP_ALIGN.RIGHT)
d._run(p, "mivada.com", 11.5, "AEAEAE", bold=True, spacing=0.04)

# ============================================================================
# 3 · AGENDA (off-white)
# ============================================================================
s = d._slide(THEME["off_white"]); head(s, "Agenda", "What we'll", "cover today.", "Agenda")
agenda = ["Why Mivada", "What we understand you need", "Workday — end to end",
          "Proven capability & outcomes", "How we deliver", "Where to from here"]
colw = CW / 2
for i, item in enumerate(agenda):
    c, r = divmod(i, 3); x = M + c * colw; y = 2.95 + r * 1.05
    _, tf = d._box(s, x, y, 0.9, 0.6); d._run(d._para(tf, first=True), f"0{i+1}", 22, CORAL, bold=True, spacing=-0.01)
    _, tf = d._box(s, x + 0.8, y + 0.04, colw - 1.0, 0.7)
    d._run(d._para(tf, first=True), item, 18, INK, bold=True, spacing=-0.01)
    d._hline(s, x, y + 0.86, colw - 0.5, THEME["hairline"], 1.0)

# ============================================================================
# 4 · WHY MIVADA — statement + differentiators (black)
# ============================================================================
s = d._slide(THEME["black"]); d._slug(s, "Why Mivada", dark=True)
d._eyebrow(s, M, 0.92, 10, "Why Mivada", color="coral")
_, tf = d._box(s, 0.92, 1.36, 11.6, 1.7); p = d._para(tf, first=True, line=0.98)
d._run(p, "We grew out of a customer.", 38, WHITE, bold=True, spacing=-0.025)
p2 = d._para(tf, line=0.98); d._run(p2, "So we ", 38, WHITE, bold=True, spacing=-0.025); d._run(p2, "speak your language.", 38, CORAL, bold=True, spacing=-0.025)
_, tf = d._box(s, 0.96, 3.12, 11, 0.5)
d._run(d._para(tf, first=True), "Your long-term partner in Workday implementations & managed services.", 15, THEME["rev_body"], bold=True, spacing=0)
diffs = [("Australian-owned, senior-led", "Not a Big 4 — practitioners on the tools"),
         ("Direct access to our CEO", "A reduced escalation matrix"),
         ("A dedicated local engagement manager", "One owner across your whole engagement"),
         ("Onshore + dedicated offshore", "Scale without losing local accountability"),
         ("10+ years, exclusively Workday", "Deep platform expertise, locally delivered")]
y = 3.95
for key, sub in diffs:
    _, tf = d._box(s, M, y, 11.4, 0.45); p = d._para(tf, first=True)
    d._run(p, "—   ", 14, CORAL, bold=True, spacing=0)
    d._run(p, key + "    ", 15, WHITE, bold=True, spacing=-0.01)
    d._run(p, sub, 12.5, THEME["rev_soft"], bold=False, spacing=0)
    y += 0.5
trusted(s, dark=True)

# ============================================================================
# 5 · WHAT WE UNDERSTAND YOU NEED (off-white) — needs as cards
# ============================================================================
s = d._slide(THEME["off_white"]); head(s, "Your needs", "What we understand", "you need.", "Your needs")
standfirst(s, "The outcomes this Workday programme has to deliver for Micah Projects.", 2.18)
needs = [("01", "Payroll compliance confidence"), ("02", "Connected HR & payroll"),
         ("03", "Rostering that enforces rules"), ("04", "A single, reliable user experience"),
         ("05", "Frontline manager visibility"), ("06", "Funding & cost visibility")]
grid(s, needs, 3, M, 2.95, CW, 1.5, 0.24, 0.24, desc=False)

# ============================================================================
# 6 · WORKDAY SERVICES — three practices (off-white)
# ============================================================================
s = d._slide(THEME["off_white"]); head(s, "What we do", "Workday,", "end to end.", "Workday services")
col = CW / 3; ytop = 2.7
for i, (name, subs) in enumerate([("People Systems & Transformation", ["HCM & core HR", "Payroll & compliance", "Time & Attendance", "HRIS implementation & optimisation"]),
                                  ("Data, Analytics & BI", ["Reporting & dashboards", "Prism & data lakehouse", "Governed analytics", "Insight at scale"]),
                                  ("Business Process & Automation", ["Process redesign", "Integrations", "Automation", "Controls & assurance"])]):
    cx = M + i * col
    if i > 0: d._vline(s, cx, ytop, 3.9, THEME["hairline"], 1.0)
    pad = 0.0 if i == 0 else 0.34
    _, tf = d._box(s, cx + pad, ytop, col - pad - 0.3, 0.9)
    d._run(d._para(tf, first=True, line=1.05), name, 18, INK, bold=True, spacing=-0.015)
    yy = ytop + 1.15
    for sub in subs:
        _, tf = d._box(s, cx + pad, yy, col - pad - 0.3, 0.5); p = d._para(tf, first=True)
        d._run(p, "—  ", 13, CORAL, bold=True, spacing=0); d._run(p, sub, 13, "323232", bold=False, spacing=0)
        yy += 0.55
_, tf = d._box(s, M, 6.62, CW, 0.4)
d._run(d._para(tf, first=True), "All wrapped in managed services & support.", 12.5, CORAL, bold=True, spacing=0.02)

# ============================================================================
# 7 · WORKDAY CAPABILITY — hero stat + proof (black)
# ============================================================================
s = d._slide(THEME["black"]); d._slug(s, "Capability", dark=True)
d._eyebrow(s, M, 1.0, 10, "Proven capability", color="coral")
_, tf = d._box(s, 0.9, 1.5, 7.2, 2.6); p = d._para(tf, first=True, line=0.82)
d._run(p, "10", 150, WHITE, bold=True, spacing=-0.04); d._run(p, "+", 150, CORAL, bold=True, spacing=-0.04)
_, tf = d._box(s, 1.0, 4.05, 6.6, 0.6)
d._run(d._para(tf, first=True), "YEARS DELIVERING WORKDAY", 14, THEME["rev_body"], bold=True, spacing=0.14, caps=True)
_, tf = d._box(s, 1.0, 4.7, 6.4, 1.6)
d._run(d._para(tf, first=True, line=1.3), "Empowering innovation through a true Workday partnership and deep local expertise — consultants average 3+ years on the platform.",
       14.5, THEME["rev_soft"], bold=False, spacing=0)
rx = 8.35; ry = 1.5
d._hline(s, rx, ry - 0.16, PAGE_W - M - rx, THEME["rev_hair"], 1.0)
for num, lab in [("30+", "Implementations & support"), ("2000+", "Workday integrations"),
                 ("3000+", "Custom reports"), ("200+", "Employees · AU, US & India"),
                 ("100+", "People-systems clients")]:
    _, tf = d._box(s, rx, ry, PAGE_W - M - rx, 0.6); p = d._para(tf, first=True, line=0.9)
    d._run(p, num, 30, WHITE, bold=True, spacing=-0.03)
    _, tf = d._box(s, rx, ry + 0.56, PAGE_W - M - rx, 0.35)
    d._run(d._para(tf, first=True), lab.upper(), 9.5, THEME["mid_grey"], bold=True, spacing=0.1, caps=True)
    ry += 1.06; d._hline(s, rx, ry - 0.16, PAGE_W - M - rx, THEME["rev_hair"], 1.0)

# ============================================================================
# 8 · CASE STUDIES (off-white)
# ============================================================================
s = d._slide(THEME["off_white"]); head(s, "Real-world success stories", "Proven Workday", "outcomes.", "Success stories")
standfirst(s, "How Mivada enables high-performance Workday operations.", 2.12, size=13)
cases = [
    ("Guzman y Gomez", "Stabilised and optimised Workday post–go-live across modules, reporting, security and compliance. Agile delivery onboarded 150+ franchisees, shipped new reporting and a redesigned security model, and lifted performance across HCM, Absence and Learning — enabling scalable, compliant global operations."),
    ("Qantas", "A partner since 2014, taking over Workday support in 2017 and becoming strategic across HR, Payroll, Digital, Automation, Data and Analytics. By scaling integrations, modernising architectures and embedding automation, Mivada grew from post–go-live support into Qantas' long-term enterprise partner."),
    ("Western Sydney Intl Airport", "Delivered a card-management system, then expanded into a strategic role on WSI's Workday program. By resolving functional and integration issues and assuming full AMS ownership, Mivada stabilised HR, Payroll, Finance and Procurement — becoming WSI's trusted long-term partner."),
]
col = CW / 3
for i, (client, para) in enumerate(cases):
    cx = M + i * col
    if i > 0: d._vline(s, cx, 3.0, 3.8, THEME["hairline"], 1.0)
    pad = 0.0 if i == 0 else 0.32
    _, tf = d._box(s, cx + pad, 2.95, col - pad - 0.25, 0.5)
    d._run(d._para(tf, first=True), client.upper(), 12, CORAL, bold=True, spacing=0.06, caps=True)
    _, tf = d._box(s, cx + pad, 3.5, col - pad - 0.28, 3.4)
    d._run(d._para(tf, first=True, line=1.28), para, 10.5, "323232", bold=False, spacing=0)

# ============================================================================
# 9 · DIVIDER — How we deliver (black)
# ============================================================================
divider("How we deliver", "Delivery you can", "see and trust.",
        "A proven methodology, a clear plan, and governance that keeps everyone aligned.")

# ============================================================================
# 10-13 · Methodology / Plan / Governance / RACI (off-white, original diagrams)
# ============================================================================
diagram("Methodology", "How we", "deliver.", 4, n=1, tag="Methodology")
# Implementation plan — native EMF
s = d._slide(THEME["off_white"]); head(s, "Indicative timeline", "Draft implementation", "plan.", "Implementation plan")
_pic = max((sh for sh in src.slides[5].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE), key=lambda x: x.width * x.height)
_emf = "/tmp/mr_emf.emf"; open(_emf, "wb").write(_pic.image.blob)
_ar2 = _pic.width / _pic.height; _w2 = 11.4; _h2 = _w2 / _ar2
if _h2 > 4.5: _h2 = 4.5; _w2 = _h2 * _ar2
s.shapes.add_picture(_emf, Inches(0.96 + (11.4 - _w2) / 2), Inches(2.45), width=Inches(_w2), height=Inches(_h2))
diagram("Governance", "How we", "run it.", 6, n=2, tag="Governance")
diagram("Roles & data", "RACI & data", "migration.", 7, n=2, tag="RACI & data migration")

# ============================================================================
# 14 · NEXT STEPS (black close)
# ============================================================================
s = d._slide(THEME["black"]); d._logo(s, "Mivada_Logo_2C_OnBlack_RGB_L", 0.96, 0.82, 0.42)
d._eyebrow(s, 0.96, 2.4, 8, "Next steps", color="coral")
_, tf = d._box(s, 0.92, 2.85, 11, 1.4); p = d._para(tf, first=True, line=0.98)
d._run(p, "Where to ", 50, WHITE, bold=True, spacing=-0.025); d._run(p, "from here?", 50, CORAL, bold=True, spacing=-0.025)
steps = ["Share your priorities and constraints", "We shape a tailored Workday plan", "Kick off with a focused discovery workshop"]
y = 4.35
for i, stp in enumerate(steps):
    _, tf = d._box(s, M, y, 0.7, 0.5); d._run(d._para(tf, first=True), f"0{i+1}", 15, CORAL, bold=True, spacing=0)
    _, tf = d._box(s, M + 0.62, y, 10, 0.5); d._run(d._para(tf, first=True), stp, 15, WHITE, bold=True, spacing=-0.01)
    y += 0.6
d._hline(s, 0.96, 6.62, CW, THEME["rev_hair"], 1.0)
_, tf = d._box(s, 0.96, 6.76, 8, 0.5); d._run(d._para(tf, first=True), "hello@mivada.com · mivada.com", 12.5, "CFCFCF", bold=True, spacing=0.02)
_, tf = d._box(s, PAGE_W - 7, 6.76, 6.04, 0.5); p = d._para(tf, first=True, align=PP_ALIGN.RIGHT)
d._run(p, "Sydney · Melbourne · AU & India", 12.5, "AEAEAE", bold=False, spacing=0.02)

OUT = "First_Meeting_with_Micah_Projects_refined.pptx"
d.save(OUT)
print(f"Wrote {OUT} ({len(d.prs.slides._sldIdLst)} slides)")
