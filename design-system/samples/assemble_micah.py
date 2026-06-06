#!/usr/bin/env python3
# ============================================================================
# Full Micah deck in Editorial, with ALL THREE versions of slides 2-4 inserted
# in sequence (Why Mivada / Workday Services / Workday Capability), so every
# option can be reviewed in context. Rest of the deck as translated.
#   python assemble_micah.py -> First_Meeting_with_Micah_Projects_variations.pptx
# ============================================================================
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "02-editorial"))
from build_pptx import Deck, THEME, PAGE_W
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image as PILImage

CORAL, INK, WHITE = "EA493F", "111111", "FFFFFF"
M = 0.96; CW = PAGE_W - 2 * M
BRANDS = "Qantas   ·   Guzman y Gomez   ·   Western Sydney International Airport"
SRC = "First_Meeting_with_Micah_Projects.pptx"
src = Presentation(SRC)
d = Deck()

# ---------- shared helpers --------------------------------------------------
def head(s, eyebrow, hpre, hacc, tag, dark=False, hsize=34, y=1.32):
    d._slug(s, tag, dark=dark)
    d._eyebrow(s, M, 0.92, 11, eyebrow, color="coral")
    _, tf = d._box(s, 0.92, y, CW, 1.2); p = d._para(tf, first=True, line=1.0)
    d._run(p, hpre + (" " if hacc else ""), hsize, (WHITE if dark else INK), bold=True, spacing=-0.022)
    if hacc: d._run(p, hacc, hsize, CORAL, bold=True, spacing=-0.022)

def standfirst(s, text, y, w=10.8, color="323232", size=14):
    _, tf = d._box(s, M, y, w, 0.9)
    d._run(d._para(tf, first=True, line=1.3), text, size, color, bold=False, spacing=0)

def trusted(s, y=6.92, dark=False):
    d._hline(s, M, y - 0.16, CW, THEME["rev_hair"] if dark else THEME["hairline"], 1.0)
    _, tf = d._box(s, M, y, CW, 0.4); p = d._para(tf, first=True)
    d._run(p, "TRUSTED BY   ", 10.5, CORAL, bold=True, spacing=0.14, caps=True)
    d._run(p, BRANDS, 10.5, THEME["rev_soft"] if dark else THEME["mid_grey"], bold=True, spacing=0.04)

def card(s, x, y, w, h, num, title, desc=None, dark=False):
    d._rect(s, x, y, w, h, fill_hex=("141414" if dark else WHITE),
            line_hex=(THEME["rev_hair"] if dark else THEME["hairline"]), line_w=1.0)
    _, tf = d._box(s, x + 0.26, y + 0.2, w - 0.45, 0.4)
    d._run(d._para(tf, first=True), num, 13.5, CORAL, bold=True, spacing=0.06)
    _, tf = d._box(s, x + 0.26, y + 0.54, w - 0.45, 0.62)
    d._run(d._para(tf, first=True, line=1.04), title, 14.5, (WHITE if dark else INK), bold=True, spacing=-0.01)
    if desc:
        _, tf = d._box(s, x + 0.26, y + 1.14, w - 0.48, h - 1.2)
        d._run(d._para(tf, first=True, line=1.16), desc, 10.5,
               (THEME["rev_soft"] if dark else "323232"), bold=False, spacing=0)

def grid(s, items, cols, x0, y0, w, rowh, gapx, gapy, dark=False, desc=True):
    cw = (w - gapx * (cols - 1)) / cols
    for i, it in enumerate(items):
        r, c = divmod(i, cols); x = x0 + c * (cw + gapx); y = y0 + r * (rowh + gapy)
        card(s, x, y, cw, rowh, it[0], it[1], (it[2] if desc and len(it) > 2 else None), dark=dark)

def imgs_from(idx, n):
    sl = src.slides[idx]
    pics = sorted([sh for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE],
                  key=lambda s: s.width * s.height, reverse=True)
    out = []
    for k, sh in enumerate(pics[:n]):
        ct = sh.image.content_type; ext = ".wmf" if "wmf" in ct else (".jpg" if "jpeg" in ct else ".png")
        raw = f"/tmp/as_{idx}_{k}{ext}"
        with open(raw, "wb") as f: f.write(sh.image.blob)
        out.append(raw)
    return out

def place_images(s, paths, x=0.96, y=2.25, w=11.4, h=4.65):
    n = len(paths);
    if n == 0: return
    gap = 0.3 if n > 1 else 0; cw = (w - gap * (n - 1)) / n; cx = x
    for p in paths:
        try: iw, ih = PILImage.open(p).size; ar = iw / ih
        except Exception: ar = 1.4
        ww = cw; hh = ww / ar
        if hh > h: hh = h; ww = hh * ar
        try: s.shapes.add_picture(p, Inches(cx + (cw - ww) / 2), Inches(y + (h - hh) / 2), width=Inches(ww))
        except Exception as e: print("  ! img skipped", e)
        cx += cw + gap

def diagram(eyebrow, hpre, hacc, idx, n=1, notes="", tag=None):
    s = d._slide(THEME["off_white"]); head(s, eyebrow, hpre, hacc, tag or eyebrow)
    place_images(s, imgs_from(idx, n))
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s

# ============================================================================
# 1 · COVER
# ============================================================================
d.cover_plain("Micah Projects · March 2026", "Implementation", "discussion.",
              "Mivada — Workday implementation & managed services", "mivada.com")

# ============================================================================
# 2-4 · WHY MIVADA (v1/v2/v3)
# ============================================================================
WHY = [
    ("01", "Australian-owned", "Creating Australian jobs, accountable locally."),
    ("02", "Not a Big 4", "Senior practitioners on the tools — no pyramid."),
    ("03", "Direct access to our CEO", "A reduced escalation matrix — decisions in hours."),
    ("04", "Dedicated engagement manager", "One local owner across your engagement."),
    ("05", "People Systems & Transformation", "HCM, Payroll, Time & Attendance and HRIS."),
    ("06", "Data, Analytics & BI", "From reporting to governed analytics."),
    ("07", "Process & Automation", "Remove the manual, error-prone work."),
    ("08", "Dedicated offshore capability", "Scale without losing local accountability."),
]
# v1
s = d._slide(THEME["off_white"]); head(s, "Why Mivada", "Your long-term", "Workday partner.", "Why Mivada · v1")
standfirst(s, "We grew out of a customer — so we understand your challenges from the inside.", 2.18)
grid(s, WHY, 4, M, 2.95, CW, 1.66, 0.22, 0.2, desc=True); trusted(s)
# v2
s = d._slide(THEME["black"]); d._slug(s, "Why Mivada · v2", dark=True)
d._eyebrow(s, M, 0.92, 10, "Why Mivada", color="coral")
_, tf = d._box(s, 0.92, 1.5, 11.6, 2.2); p = d._para(tf, first=True, line=0.98)
d._run(p, "We grew out of a customer.", 44, WHITE, bold=True, spacing=-0.025)
p2 = d._para(tf, line=0.98)
d._run(p2, "So we ", 44, WHITE, bold=True, spacing=-0.025); d._run(p2, "get it.", 44, CORAL, bold=True, spacing=-0.025)
_, tf = d._box(s, 0.96, 3.55, 11, 0.5)
d._run(d._para(tf, first=True), "Your long-term partner in Workday implementations & managed services.",
       15, THEME["rev_body"], bold=True, spacing=0)
for ci, group in enumerate([WHY[:4], WHY[4:]]):
    cx = M + ci * (CW / 2)
    for ri, it in enumerate(group):
        ry = 4.35 + ri * 0.62
        _, tf = d._box(s, cx, ry, 0.7, 0.5); d._run(d._para(tf, first=True), it[0], 14, CORAL, bold=True, spacing=0)
        _, tf = d._box(s, cx + 0.62, ry, CW / 2 - 0.9, 0.5)
        d._run(d._para(tf, first=True), it[1], 14, WHITE, bold=True, spacing=-0.01)
trusted(s, dark=True)
# v3
s = d._slide(THEME["off_white"]); d._slug(s, "Why Mivada · v3")
d._eyebrow(s, M, 1.15, 5, "Why Mivada", color="coral")
_, tf = d._box(s, 0.92, 1.6, 5.3, 1.9); p = d._para(tf, first=True, line=1.0)
d._run(p, "Built by people who ", 34, INK, bold=True, spacing=-0.022); d._run(p, "ran it.", 34, CORAL, bold=True, spacing=-0.022)
_, tf = d._box(s, M, 3.65, 5.0, 2.6)
d._run(d._para(tf, first=True, line=1.34),
       "We grew out of a customer, so we understand your challenges first-hand. Australian-owned and "
       "senior-led — not a Big 4 — with direct access to our CEO and a dedicated local engagement "
       "manager who owns your outcome.", 13.5, "323232", bold=False, spacing=0)
rx, rw = 6.95, PAGE_W - M - 6.95; y = 1.3; rh = 1.12
for k, sub in [("People Systems & Transformation", "HCM · Payroll · T&A · HRIS"),
               ("Data, Analytics & BI", "Reporting · dashboards · governed analytics"),
               ("Process & Automation", "Redesign · integrations · automation"),
               ("Dedicated offshore capability", "Scale with local accountability")]:
    d._hline(s, rx, y, rw, THEME["hairline"], 1.0)
    _, tf = d._box(s, rx, y + 0.2, rw, 0.5); d._run(d._para(tf, first=True), k, 16, INK, bold=True, spacing=-0.01)
    _, tf = d._box(s, rx, y + 0.62, rw, 0.4); d._run(d._para(tf, first=True), sub, 11.5, CORAL, bold=True, spacing=0.02)
    y += rh
d._hline(s, rx, y, rw, THEME["hairline"], 1.0); trusted(s)

# ============================================================================
# 5-7 · WORKDAY SERVICES (v1/v2/v3)
# ============================================================================
SERVICES9 = [("01", "Advisory & roadmap"), ("02", "Implementation & deployment"), ("03", "Optimisation & phase 2"),
             ("04", "Managed services (AMS)"), ("05", "Integrations"), ("06", "Data, analytics & BI"),
             ("07", "Reporting & dashboards"), ("08", "Security & compliance"), ("09", "Automation")]
# v1
s = d._slide(THEME["off_white"]); head(s, "What we do", "Workday,", "end to end.", "Workday services · v1")
standfirst(s, "One partner across the whole platform — from first roadmap to long-term managed services.", 2.18)
grid(s, SERVICES9, 3, M, 2.95, CW, 1.16, 0.22, 0.2, desc=False)
# v2
s = d._slide(THEME["black"]); head(s, "Workday services", "Across the whole", "Workday lifecycle.", "Workday services · v2", dark=True)
col = CW / 4; ytop = 3.4
for i, (name, sub) in enumerate([("Advise", "Roadmap, business case, design"),
                                 ("Implement", "Deploy HCM, Payroll, Finance & more"),
                                 ("Optimise", "Phase 2, integrations, reporting"),
                                 ("Manage", "AMS, automation, continuous value")]):
    cx = M + i * col
    if i > 0: d._vline(s, cx, ytop, 2.4, THEME["rev_hair"], 1.0)
    pad = 0.0 if i == 0 else 0.34
    _, tf = d._box(s, cx + pad, ytop, col - pad - 0.25, 0.5)
    d._run(d._para(tf, first=True), f"STAGE 0{i+1}", 11, CORAL, bold=True, spacing=0.16, caps=True)
    _, tf = d._box(s, cx + pad, ytop + 0.5, col - pad - 0.25, 0.7)
    d._run(d._para(tf, first=True), name, 26, WHITE, bold=True, spacing=-0.025)
    _, tf = d._box(s, cx + pad, ytop + 1.25, col - pad - 0.3, 1.2)
    d._run(d._para(tf, first=True, line=1.25), sub, 12, THEME["rev_soft"], bold=False, spacing=0)
_, tf = d._box(s, M, 6.5, CW, 0.4)
d._run(d._para(tf, first=True), "Wrapped in managed services & support at every stage.", 12.5, THEME["mid_grey"], bold=True, spacing=0.02)
# v3
s = d._slide(THEME["off_white"]); head(s, "Services", "Three practices,", "one platform.", "Workday services · v3")
col = CW / 3; ytop = 2.7
for i, (name, subs) in enumerate([("People Systems & Transformation", ["HCM & core HR", "Payroll & compliance", "Time & Attendance", "HRIS implementation & optimisation"]),
                                  ("Data, Analytics & BI", ["Reporting & dashboards", "Prism & data lakehouse", "Governed analytics", "Insight at scale"]),
                                  ("Business Process & Automation", ["Process redesign", "Integrations", "Automation", "Controls & assurance"])]):
    cx = M + i * col
    if i > 0: d._vline(s, cx, ytop, 3.9, THEME["hairline"], 1.0)
    pad = 0.0 if i == 0 else 0.34
    _, tf = d._box(s, cx + pad, ytop, col - pad - 0.3, 0.9)
    d._run(d._para(tf, first=True, line=1.05), name, 18, INK, bold=True, spacing=-0.015)
    yy = ytop + 1.15
    for sub in subs:
        _, tf = d._box(s, cx + pad, yy, col - pad - 0.3, 0.5); p = d._para(tf, first=True)
        d._run(p, "—  ", 13, CORAL, bold=True, spacing=0); d._run(p, sub, 13, "323232", bold=False, spacing=0)
        yy += 0.55
_, tf = d._box(s, M, 6.62, CW, 0.4)
d._run(d._para(tf, first=True), "All wrapped in managed services & support.", 12.5, CORAL, bold=True, spacing=0.02)

# ============================================================================
# 8-10 · WORKDAY CAPABILITY (v1/v2/v3)
# ============================================================================
STATS = [("3", "+", "Certifications / consultant"), ("2000", "+", "Workday integrations"),
         ("3000", "+", "Workday custom reports"), ("30", "+", "Implementations & support"),
         ("200", "+", "Employees · AU, US & India"), ("100", "+", "People-systems clients"),
         ("10", "+", "Years delivering Workday")]
# v1
s = d._slide(THEME["off_white"]); head(s, "Capability", "Capability in", "Workday.", "Capability · v1")
standfirst(s, "Empowering innovation through a Workday partnership and local expertise — consultants average "
              "3+ years' Workday experience.", 2.18)
d._stat_row(s, 3.7, STATS[:4], light=True); d._stat_row(s, 5.6, STATS[4:], light=True)
# v2
s = d._slide(THEME["black"]); d._slug(s, "Capability · v2", dark=True)
d._eyebrow(s, M, 1.0, 10, "Proven capability", color="coral")
_, tf = d._box(s, 0.9, 1.5, 7.2, 2.6); p = d._para(tf, first=True, line=0.82)
d._run(p, "10", 150, WHITE, bold=True, spacing=-0.04); d._run(p, "+", 150, CORAL, bold=True, spacing=-0.04)
_, tf = d._box(s, 1.0, 4.05, 6.6, 0.6)
d._run(d._para(tf, first=True), "YEARS DELIVERING WORKDAY", 14, THEME["rev_body"], bold=True, spacing=0.14, caps=True)
_, tf = d._box(s, 1.0, 4.7, 6.4, 1.6)
d._run(d._para(tf, first=True, line=1.3), "Empowering innovation through a true Workday partnership and deep local expertise.",
       15, THEME["rev_soft"], bold=False, spacing=0)
rx = 8.35; ry = 1.6
d._hline(s, rx, ry - 0.16, PAGE_W - M - rx, THEME["rev_hair"], 1.0)
for num, lab in [("30+", "Implementations & support"), ("200+", "Employees · AU, US & India"),
                 ("2000+", "Workday integrations"), ("3000+", "Custom reports")]:
    _, tf = d._box(s, rx, ry, PAGE_W - M - rx, 0.7); p = d._para(tf, first=True, line=0.9)
    d._run(p, num, 34, WHITE, bold=True, spacing=-0.03)
    _, tf = d._box(s, rx, ry + 0.66, PAGE_W - M - rx, 0.4)
    d._run(d._para(tf, first=True), lab.upper(), 10, THEME["mid_grey"], bold=True, spacing=0.1, caps=True)
    ry += 1.25; d._hline(s, rx, ry - 0.18, PAGE_W - M - rx, THEME["rev_hair"], 1.0)
# v3
s = d._slide(THEME["off_white"]); head(s, "Capability", "The numbers behind", "the partnership.", "Capability · v3")
cards = [("3+", "Certifications / consultant"), ("2000+", "Workday integrations"), ("3000+", "Custom reports"),
         ("30+", "Implementations & support"), ("200+", "Employees · AU, US & India"), ("100+", "People-systems clients")]
cols = 3; gx, gy = 0.24, 0.24; cw = (CW - gx * (cols - 1)) / cols; ch = 1.55; y0 = 2.95
for i, (num, lab) in enumerate(cards):
    r, c = divmod(i, cols); x = M + c * (cw + gx); y = y0 + r * (ch + gy)
    d._rect(s, x, y, cw, ch, fill_hex=WHITE, line_hex=THEME["hairline"], line_w=1.0)
    _, tf = d._box(s, x + 0.3, y + 0.26, cw - 0.5, 0.8); p = d._para(tf, first=True, line=0.9)
    d._run(p, num.rstrip("+"), 40, INK, bold=True, spacing=-0.035)
    d._run(p, "+" if num.endswith("+") else "", 40, CORAL, bold=True, spacing=-0.035)
    _, tf = d._box(s, x + 0.3, y + 1.04, cw - 0.5, 0.45)
    d._run(d._para(tf, first=True), lab.upper(), 10.5, THEME["mid_grey"], bold=True, spacing=0.1, caps=True)
_, tf = d._box(s, M, y0 + 2 * (ch + gy) + 0.05, CW, 0.5)
d._run(d._para(tf, first=True), "10+ years delivering Workday · consultants average 3+ years' experience.",
       13, CORAL, bold=True, spacing=0.02)

# ============================================================================
# 11+ · rest of the deck (as translated)
# ============================================================================
diagram("How we deliver", "Our", "methodology.", 4, n=1, tag="Methodology")
# Draft implementation plan — native EMF (renders in PowerPoint)
s = d._slide(THEME["off_white"]); head(s, "Indicative timeline", "Draft implementation", "plan.", "Implementation plan")
_pic = max((sh for sh in src.slides[5].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE), key=lambda x: x.width * x.height)
_emf = "/tmp/as_emf.emf"; open(_emf, "wb").write(_pic.image.blob)
_ar = _pic.width / _pic.height; _w = 11.4; _h = _w / _ar
if _h > 4.5: _h = 4.5; _w = _h * _ar
s.shapes.add_picture(_emf, Inches(0.96 + (11.4 - _w) / 2), Inches(2.45), width=Inches(_w), height=Inches(_h))
diagram("How we run it", "", "Governance.", 6, n=2, tag="Governance")
diagram("Roles & data", "RACI & data", "migration.", 7, n=2, tag="RACI & data migration")
d.content("Understanding your needs", "What you need", "this to do.",
          ["Payroll compliance confidence", "Connected HR & payroll", "Rostering that enforces rules",
           "A single, reliable user experience", "Frontline manager visibility", "Funding & cost visibility"])
# Success stories (3 columns)
s = d._slide(THEME["off_white"]); head(s, "Real-world success stories", "Proven Workday", "outcomes.", "Success stories")
standfirst(s, "How Mivada enables high-performance Workday operations.", 2.12, size=13)
cases = [
    ("Guzman y Gomez", "Mivada partnered with GYG to stabilise and optimise their Workday platform post–go-live across modules, reporting, security and compliance. Through Agile delivery we onboarded 150+ franchisees, delivered new reporting and a redesigned security model, and improved performance across HCM, Absence and Learning — enabling scalable, compliant global operations."),
    ("Qantas", "Mivada has supported Qantas since 2014, taking over Workday support in 2017 and rapidly becoming a strategic partner across HR, Payroll, Digital, Automation, Data and Analytics. By scaling integrations, modernising architectures and embedding automation, Mivada evolved from post–go-live support to Qantas' long-term enterprise partner."),
    ("Western Sydney Intl Airport", "Mivada partnered with WSI to deliver a card-management system, then expanded into a strategic role supporting WSI's Workday program. By resolving functional and integration issues and assuming full AMS ownership, Mivada stabilised HR, Payroll, Finance and Procurement — becoming WSI's trusted long-term partner."),
]
col = CW / 3
for i, (client, para) in enumerate(cases):
    cx = M + i * col
    if i > 0: d._vline(s, cx, 3.0, 3.8, THEME["hairline"], 1.0)
    pad = 0.0 if i == 0 else 0.32
    _, tf = d._box(s, cx + pad, 2.95, col - pad - 0.25, 0.5)
    d._run(d._para(tf, first=True), client.upper(), 12, CORAL, bold=True, spacing=0.06, caps=True)
    _, tf = d._box(s, cx + pad, 3.5, col - pad - 0.28, 3.4)
    d._run(d._para(tf, first=True, line=1.26), para, 10, "323232", bold=False, spacing=0)
# Questions
d.contact_plain("Over to you", "Questions?", "", "hello@mivada.com · mivada.com", "Sydney · Melbourne · AU & India")

OUT = "First_Meeting_with_Micah_Projects_variations.pptx"
d.save(OUT)
print(f"Wrote {OUT} ({len(d.prs.slides._sldIdLst)} slides)")
