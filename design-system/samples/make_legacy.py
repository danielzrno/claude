#!/usr/bin/env python3
# ============================================================================
# STAND-IN "before" deck — a generic, off-brand company overview, the kind of
# existing PowerPoint you'd be handed to restyle. Default Office look (Calibri,
# a flat blue accent, bullet lists, a table, a bar chart, a quote). Used only to
# demonstrate the Editorial translation workflow until the real deck is provided.
#
#   python make_legacy.py   ->  legacy-overview.pptx
# ============================================================================
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

BLUE = RGBColor(0x1F, 0x4E, 0x79)      # generic corporate blue
GREY = RGBColor(0x59, 0x59, 0x59)
LGREY = RGBColor(0xD9, 0xD9, 0xD9)
BLACK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5

def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    return tf

def run(p, text, size, color=BLACK, bold=False, italic=False):
    r = p.add_run(); r.text = text; f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = FONT; f.color.rgb = color
    return r

def para(tf, first=False, align=PP_ALIGN.LEFT, space=6, level=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space); p.level = level
    return p

def title_bar(s, text, kicker=None):
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.6)
    p = tf.paragraphs[0]; run(p, text, 26, WHITE, bold=True)
    if kicker:
        tf2 = box(s, 0.6, 1.35, 12, 0.5)
        run(para(tf2, first=True), kicker, 13, GREY, italic=True)

def bullets(s, items, x=0.7, y=1.9, w=12, size=16):
    tf = box(s, x, y, w, H - y - 0.6)
    for i, it in enumerate(items):
        lvl = 0
        if isinstance(it, tuple): it, lvl = it
        p = para(tf, first=(i == 0), space=10, level=lvl)
        run(p, ("•  " if lvl == 0 else "–  ") + it, size, BLACK)

def footer(s, n):
    tf = box(s, 0.6, H - 0.5, 12.1, 0.4)
    p = para(tf, first=True); run(p, "Mivada  |  Company Overview  |  Confidential", 9, GREY)
    p2 = box(s, W - 1.3, H - 0.5, 0.8, 0.4); pr = para(p2, first=True, align=PP_ALIGN.RIGHT)
    run(pr, str(n), 9, GREY)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

# (1) Title
s = slide()
tf = box(s, 0.8, 2.5, 11.7, 2.5)
run(para(tf, first=True), "Mivada", 54, BLUE, bold=True)
run(para(tf, space=4), "Company Overview", 30, BLACK)
run(para(tf), "Capability Statement — 2026", 16, GREY)
notes(s, "Intro the company. Australian technology consultancy. Thank the audience for their time.")

# (2) About us  (with an embedded image, to exercise image-carrying on translation)
def _placeholder_photo(path, w=1200, h=820):
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (w, h), (208, 212, 209)); d = ImageDraw.Draw(img)
    for y in range(h):
        t = y / h; d.line([(0, y), (w, y)], fill=(int(206 - 34 * t), int(210 - 30 * t), int(207 - 28 * t)))
    d.ellipse([w * 0.52, h * 0.18, w * 0.96, h * 0.86], fill=(156, 165, 162))
    d.rectangle([w * 0.05, h * 0.6, w * 0.42, h * 0.94], fill=(126, 134, 131))
    img.save(path)
PHOTO = "/tmp/_legacy_photo.png"; _placeholder_photo(PHOTO)

s = slide(); title_bar(s, "About Us")
tf = box(s, 0.7, 1.5, 7.3, 1.8)
run(para(tf, first=True),
    "Mivada is an Australian technology consultancy founded in 2014. We turn enterprise "
    "platforms — Workday, payroll, data and AI — into measurable outcomes, built around how "
    "teams actually work. Onshore in Australia with a delivery team across AU and India.", 15)
bullets(s, ["Founded 2014; formerly LJM Infotech",
            "120+ certified specialists",
            "Delivery across Australia and India",
            "Independent and platform-agnostic"], x=0.7, y=3.4, w=7.0, size=15)
s.shapes.add_picture(PHOTO, Inches(8.2), Inches(1.7), width=Inches(4.5))
footer(s, 2); notes(s, "People first. We are not a body shop — outcomes, not hours.")

# (3) What we do
s = slide(); title_bar(s, "What We Do", "Four service pillars")
bullets(s, ["Workday & ERP — implementation, optimisation and managed support",
            "Payroll consulting — compliant, accurate payroll at scale",
            "Data & AI — lakehouse architecture, governed analytics, applied AI",
            "Intelligent automation — remove manual, error-prone work"], y=2.0, size=18)
footer(s, 3); notes(s, "Lead with Workday & ERP — our anchor practice.")

# (4) Our approach
s = slide(); title_bar(s, "Our Approach", "How we work")
bullets(s, ["1. Listen — understand the people and the work before the technology",
            "2. Design — map to real-world needs",
            "3. Deliver — implement in increments; value early",
            "4. Care — measure outcomes; keep improving after go-live"], y=2.0, size=18)
footer(s, 4); notes(s, "Human-first, end to end. The 'Care' phase is our differentiator.")

# (5) By the numbers — TABLE
s = slide(); title_bar(s, "By the Numbers")
rows = [("Metric", "Value"),
        ("Founded", "2014"),
        ("Certified specialists", "120+"),
        ("Implementations delivered", "40+"),
        ("Client retention", "98%"),
        ("Delivery footprint", "AU + India")]
tbl_shape = s.shapes.add_table(len(rows), 2, Inches(0.7), Inches(1.7), Inches(7.5), Inches(4.2))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4.8); tbl.columns[1].width = Inches(2.7)
for r, (a, b) in enumerate(rows):
    for c, val in enumerate((a, b)):
        cell = tbl.cell(r, c); cell.text = val
        pr = cell.text_frame.paragraphs[0]; pr.runs[0].font.name = FONT
        pr.runs[0].font.size = Pt(14)
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
            pr.runs[0].font.color.rgb = WHITE; pr.runs[0].font.bold = True
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if r % 2 else LGREY
footer(s, 5); notes(s, "Figures are illustrative samples. Replace with verified numbers.")

# (6) Results — BAR CHART
s = slide(); title_bar(s, "Results", "Projects delivered per year")
cd = CategoryChartData()
cd.categories = ["2022", "2023", "2024", "2025"]
cd.add_series("Projects delivered", (12, 19, 28, 41))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                        Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.6), cd)
chart = gf.chart; chart.has_legend = False
footer(s, 6); notes(s, "Steady growth. 41 projects in 2025. All figures illustrative.")

# (7) Testimonial
s = slide(); title_bar(s, "What Clients Say")
tf = box(s, 1.0, 2.4, 11.3, 3, anchor=MSO_ANCHOR.MIDDLE)
run(para(tf, first=True),
    "“They delivered in weeks what we’d scoped for a year — and our team actually "
    "uses it.”", 26, BLACK, italic=True)
run(para(tf, space=0), "— Operations Director, enterprise client", 15, GREY)
footer(s, 7); notes(s, "Strongest proof point. Reinforces the 'Care'/adoption message.")

# (8) Contact
s = slide()
bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(H))
bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background(); bar.shadow.inherit = False
tf = bar.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.8)
run(para(tf, first=True), "Let’s talk", 40, WHITE, bold=True)
run(para(tf, space=2), "hello@mivada.com   |   mivada.com", 18, WHITE)
run(para(tf), "Sydney & Melbourne   |   Onshore AU & India", 14, RGBColor(0xCF, 0xDA, 0xE6))
notes(s, "Close with the invitation. Offer a discovery workshop.")

prs.save("legacy-overview.pptx")
print(f"Wrote legacy-overview.pptx ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
