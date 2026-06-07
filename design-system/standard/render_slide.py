#!/usr/bin/env python3
"""Faithful-enough Pillow renderer for a single PPTX slide (layout + per-run colour)."""
import sys, re
from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

import os as _os
EMU = 914400; DPI = 96; FRAME = True
_FD = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "assets", "fonts")
def _font_path(name, fallback):
    for p in (_os.path.join(_FD, name), "/tmp/" + name, fallback):
        if _os.path.exists(p): return p
    return fallback
REG = _font_path("Inter-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")
BLD = _font_path("Inter-Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")
_fc = {}
def font(sz, bold):
    key = (round(sz * 2) / 2, bold)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(BLD if bold else REG, max(6, round(sz * DPI / 72)))
    return _fc[key]
def px(emu): return emu / EMU * DPI
def hexof(c):
    try:
        if c.type is not None: return "#" + str(c.rgb)
    except Exception: pass
    return None
def sfill(sh):
    try:
        if sh.fill.type == 1: return "#" + str(sh.fill.fore_color.rgb)
    except Exception: pass
    return None
def sline(sh):
    try:
        return "#" + str(sh.line.color.rgb), max(1, round((sh.line.width.pt if sh.line.width else 1) * DPI / 72))
    except Exception: return None, 0
def bg_color(slide):
    s = etree.tostring(slide.background.element).decode()
    m = re.search(r"<p:bg>.*?</p:bg>", s, re.S)
    if m:
        sr = re.search(r'srgbClr val="(\w{6})"', m.group(0))
        if sr: return "#" + sr.group(1)
        sc = re.search(r'schemeClr val="(\w+)"', m.group(0))
        if sc and sc.group(1) in ("tx1", "dk1"): return "#000000"
    return "#FFFFFF"

def layout_para(d, p, w, default_col):
    pad = 4
    toks = []
    for r in p.runs:
        sz = r.font.size.pt if r.font.size else 12
        f = font(sz, bool(r.font.bold)); col = hexof(r.font.color) or default_col
        for seg in re.split(r"([\n\x0b])", r.text):
            if seg in ("\n", "\x0b"):
                toks.append(("__BR__", f, col, sz)); continue
            for k, wd in enumerate(seg.split(" ")):
                if k > 0: toks.append((" ", f, col, sz))
                if wd: toks.append((wd, f, col, sz))
    maxw = w - 2 * pad
    lines, cur, curw = [], [], 0
    for wd, f, col, sz in toks:
        if wd == "__BR__":
            lines.append(cur); cur, curw = [], 0; continue
        ww = d.textlength(wd, font=f)
        if wd == " ":
            if cur: cur.append((wd, f, col, sz, ww)); curw += ww
            continue
        if curw + ww > maxw and cur:
            lines.append(cur); cur, curw = [], 0
        cur.append((wd, f, col, sz, ww)); curw += ww
    if cur: lines.append(cur)
    out = [(ln, max((t[3] for t in ln), default=12) * DPI / 72 * 1.22, p.alignment) for ln in lines]
    return out

def draw_frame(d, tf, x, y, w, h, default_col):
    pad = 4
    all_lines = []
    for p in tf.paragraphs:
        if any(r.text for r in p.runs):
            all_lines += layout_para(d, p, w, default_col)
    if not all_lines: return
    total = sum(lh for _, lh, _ in all_lines)
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE: cy = y + (h - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM: cy = y + h - total - pad
    else: cy = y + pad
    for ln, lh, align in all_lines:
        lw = sum(t[4] for t in ln)
        if align == PP_ALIGN.CENTER: cx = x + (w - lw) / 2
        elif align == PP_ALIGN.RIGHT: cx = x + w - pad - lw
        else: cx = x + pad
        for wd, f, col, sz, ww in ln:
            d.text((cx, cy + (lh - sz * DPI / 72 * 1.22)), wd, font=f, fill=col)
            cx += ww
        cy += lh
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
def group_tf(g):
    xf = g._element.find(f".//{A_NS}xfrm")
    off = xf.find(f"{A_NS}off"); ext = xf.find(f"{A_NS}ext")
    cof = xf.find(f"{A_NS}chOff"); cex = xf.find(f"{A_NS}chExt")
    ox, oy = int(off.get("x")), int(off.get("y"))
    ew, eh = int(ext.get("cx")), int(ext.get("cy"))
    cx0, cy0 = int(cof.get("x")), int(cof.get("y"))
    cw0, ch0 = int(cex.get("cx")) or 1, int(cex.get("cy")) or 1
    sx, sy = ew / cw0, eh / ch0
    return lambda x, y, w, h: (ox + (x - cx0) * sx, oy + (y - cy0) * sy, w * sx, h * sy)

def flatten(shapes, tf):
    for sh in shapes:
        if str(sh.shape_type).startswith("GROUP"):
            inner = group_tf(sh)
            yield from flatten(sh.shapes, lambda x, y, w, h, o=tf, i=inner: o(*i(x, y, w, h)))
        elif sh.left is not None:
            yield sh, tf(sh.left, sh.top, sh.width, sh.height)

def render(path, idx, out):
    prs = Presentation(path)
    W, H = round(px(prs.slide_width)), round(px(prs.slide_height))
    slide = prs.slides[idx]
    bg = bg_color(slide)
    img = Image.new("RGB", (W, H), bg); d = ImageDraw.Draw(img)
    default_text = "#111111" if bg.upper() != "#000000" else "#FFFFFF"
    for sh, (ax, ay, aw, ah) in flatten(slide.shapes, lambda x, y, w, h: (x, y, w, h)):
        x, y, w, h = px(ax), px(ay), px(aw), px(ah)
        tag = sh._element.tag
        if tag.endswith("}cxnSp") or str(sh.shape_type).startswith("LINE"):
            lc, lw = sline(sh)
            xf = sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            flipV = xf is not None and xf.get("flipV") == "1"
            p1, p2 = ((x, y), (x + w, y + h)) if not flipV else ((x, y + h), (x + w, y))
            d.line([p1, p2], fill=lc or "#AEAEAE", width=lw or 1); continue
        if sh.shape_type == 1:
            fill = sfill(sh); lc, lw = sline(sh)
            d.rectangle([x, y, x + w, y + h], fill=fill, outline=lc, width=lw or 1)
        elif str(sh.shape_type).startswith("PICTURE"):
            try:
                from io import BytesIO
                im = Image.open(BytesIO(sh.image.blob)).convert("RGBA")
                im.thumbnail((round(w), round(h)))
                img.paste(im, (round(x), round(y)), im)
            except Exception:
                d.rectangle([x, y, x + w, y + h], outline="#AAAAAA")
        if sh.has_text_frame and sh.text_frame.text.strip():
            draw_frame(d, sh.text_frame, x, y, w, h, default_text)
    if FRAME:
        d.rectangle([0, 0, W - 1, H - 1], outline="#EA493F", width=2)
    img.save(out); print("wrote", out)

if __name__ == "__main__":
    render(sys.argv[1], int(sys.argv[2]), sys.argv[3])
