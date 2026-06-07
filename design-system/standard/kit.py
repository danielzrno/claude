#!/usr/bin/env python3
"""Mivada Standard builder kit — primitives for composing NEW on-style slides.

The box idiom (match slides 'Four ways' / 'Workday lifecycle'): a filled/bordered
rectangle with a CORAL LABEL inside — never a left-edge accent stripe, never a
drop shadow. Coral is the single accent. Build slides light or dark; alternate.

    from kit import Deck, CORAL, INK, GRAPH
    d = Deck()                       # 16:9, logo + theme handled per slide
    s = d.slide()                    # auto-alternates light/dark each call
    d.eyebrow(s, "DELIVERY APPROACH"); d.title(s, "Velocity & value — ", "Workday GO.")
    d.lead(s, "Workday's packaged deployment — live fast.")
    d.cards(s, 2.4, [("90%", "Pre-configured", "…"), ("10%", "Unique to you", "…")])
    d.flow(s, 4.8, [("1. Mobilise","Charter"),("2. Architect","Workshops"), …])
    d.save("MyDeck.pptx")            # then theme_deck.build() for B&W, build_web.build() for HTML/Word
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

def H(s): return RGBColor.from_string(s)
CORAL, INK, GRAPH = H("EA493F"), H("111111"), H("323232")
WHITE, OFFW, BORDER = H("FFFFFF"), H("FAFAFA"), H("E4E4E0")
BLACK, DCARD, DBORDER, DBODY = H("000000"), H("141414"), H("3A3A3A"), H("CFCFCF")
LINEC, FONT = H("AEAEAE"), "Inter"
HERE = os.path.dirname(os.path.abspath(__file__))

def _theme(dark):
    return dict(bg=BLACK if dark else OFFW, card=DCARD if dark else WHITE,
                border=DBORDER if dark else BORDER, label=CORAL,
                head=WHITE if dark else INK, body=DBODY if dark else GRAPH,
                line=DBORDER if dark else LINEC)

class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self._dark = True            # first slide dark (cover/opener); then alternates
        self.logo = None
        m = os.path.join(HERE, "Mivada_Standard.pptx")
        if os.path.exists(m):
            for s in Presentation(m).slides:
                for sh in s.shapes:
                    if str(sh.shape_type).startswith("PICTURE") and sh.width and sh.width < Inches(1):
                        self.logo = sh.image.blob; break
                if self.logo: break

    def slide(self, dark=None):
        if dark is None: dark = self._dark
        self._dark = not dark
        sl = self.prs.slides.add_slide(self.blank)
        t = _theme(dark); sl._t = t
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb = t["bg"]
        icon = "Mivada_Icon_White_RGB_L.png" if dark else "Mivada_Icon_Melon_RGB_L.png"
        p = os.path.join(HERE, "..", "assets", "logos", icon)
        if os.path.exists(p): sl.shapes.add_picture(p, Inches(0.96), Inches(0.5), Inches(0.44), Inches(0.26))
        return sl

    # ---- primitives ----
    def rect(self, s, l, t, w, h, fill=None, line=None, lw=1.0):
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        sp.shadow.inherit = False
        if fill is None: sp.fill.background()
        else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None: sp.line.fill.background()
        else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
        sp.text_frame.margin_left = sp.text_frame.margin_right = 0
        return sp

    def text(self, s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=None):
        tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame
        tf.word_wrap = True
        for m in ("left", "right", "top", "bottom"): setattr(tf, f"margin_{m}", 0)
        tf.vertical_anchor = anchor
        for i, para in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
            if ls: p.line_spacing = ls
            for txt, size, bold, col in para:
                r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold
                r.font.color.rgb = col; r.font.name = FONT
        return tb

    def line(self, s, x1, y1, x2, y2, color=None, w=1.0):
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = color or s._t["line"]; cn.line.width = Pt(w); cn.shadow.inherit = False

    def eyebrow(self, s, txt):
        self.text(s, 6.5, 0.46, 6.0, 0.35, [[(txt, 11, True, CORAL)]], align=PP_ALIGN.RIGHT)
    def title(self, s, ink_part, coral_part):
        self.text(s, 0.92, 0.98, 11.41, 0.95, [[(ink_part, 33, True, s._t["head"]), (coral_part, 33, True, CORAL)]])
    def lead(self, s, txt):
        self.text(s, 0.96, 1.74, 11.42, 0.5, [[(txt, 12.5, False, s._t["body"])]], ls=1.12)

    # ---- the standard box: fill + border + coral label inside, no stripe ----
    def card(self, s, l, t, w, h, label, head=None, body=None):
        T = s._t
        self.rect(s, l, t, w, h, fill=T["card"], line=T["border"])
        y = t + 0.2
        if label: self.text(s, l + 0.3, y, w - 0.6, 0.3, [[(label, 11.5, True, CORAL)]]); y += 0.34
        if head:  self.text(s, l + 0.3, y, w - 0.6, 0.32, [[(head, 14, True, T["head"])]]); y += 0.4
        if body:  self.text(s, l + 0.3, y, w - 0.6, h - (y - t) - 0.2, [[(body, 11, False, T["body"])]], ls=1.12)

    def cards(self, s, top, items, h=1.7, left=0.96, right=12.38, gap=0.4):
        """Row of evenly-spaced standard boxes: items=[(label, head, body)]."""
        n = len(items); w = (right - left - gap * (n - 1)) / n
        for i, it in enumerate(items):
            self.card(s, left + i * (w + gap), top, w, h, *((list(it) + [None, None])[:3]))

    def flow(self, s, top, items, h=1.2, emph=(), left=0.96, right=12.38):
        """Cards joined by coral arrows: items=[(name, detail)]. emph=coral-filled indices."""
        T = s._t; n = len(items); gap = 0.46 if n <= 5 else 0.34; w = (right - left - gap * (n - 1)) / n
        for i, (name, detail) in enumerate(items):
            x = left + i * (w + gap); on = i in emph
            self.rect(s, x, top, w, h, fill=(CORAL if on else T["card"]), line=(None if on else T["border"]))
            para = [[(name, 11, True, WHITE if on else CORAL)]]
            if detail: para.append([(detail, 9, False, WHITE if on else T["body"])])
            self.text(s, x + 0.16, top, w - 0.32, h, para, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ls=1.05)
            if i < n - 1:
                self.text(s, x + w + gap / 2 - 0.25, top + h / 2 - 0.25, 0.5, 0.5,
                          [[("→", 18, True, CORAL)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def band(self, s, top, lead, rest, h=0.9, left=0.96, w=11.42):
        """Full-width black emphasis band: coral lead + body."""
        self.rect(s, left, top, w, h, fill=INK)
        self.text(s, left + 0.36, top, w - 0.72, h, [[(lead + "  ", 12, True, CORAL), (rest, 11, False, WHITE)]],
                  anchor=MSO_ANCHOR.MIDDLE, ls=1.12)

    def grid(self, s, top, cols, rows, header=True, left=0.96, rh=0.5):
        """Boxed table. cols=[(name,width)]; rows=[[cell,…]]. First row coral header if header."""
        T = s._t; y = top
        if header:
            x = left
            for name, w in cols:
                self.rect(s, x, y, w, 0.42, fill=CORAL); self.text(s, x + 0.14, y, w - 0.28, 0.42,
                          [[(name, 10, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE); x += w
            y += 0.48
        for row in rows:
            x = left
            for (name, w), val in zip(cols, row):
                self.rect(s, x, y, w, rh, fill=T["card"], line=T["border"])
                self.text(s, x + 0.14, y, w - 0.28, rh, [[(val, 9.5, False, T["body"])]], anchor=MSO_ANCHOR.MIDDLE, ls=1.05)
                x += w
            y += rh + 0.05

    def save(self, path): self.prs.save(path); return path
