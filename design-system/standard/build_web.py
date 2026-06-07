#!/usr/bin/env python3
"""Generate web/print artifacts from a pair of themed decks (light + dark).
Renders every slide (hi-res, Inter) and assembles an interactive HTML deck
(light/dark toggle, keyboard nav, print 1/page), A4 landscape/portrait docs,
and Word (.docx) one full-page slide per page — for both themes.

  python build_web.py                      # rebuild the Standard's artifacts
  python build_web.py --light X_Light.pptx --dark X_Dark.pptx --basename X --title "X"
"""
import os, importlib.util, shutil
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("rs", os.path.join(os.path.dirname(__file__), "render_slide.py"))
if not os.path.exists(spec.origin):                      # fall back to the session copy
    spec = importlib.util.spec_from_file_location("rs", "/tmp/render_slide.py")
rs = importlib.util.module_from_spec(spec); spec.loader.exec_module(rs)
rs.DPI = 144            # 1.5x → 1920x1080
rs.FRAME = False
rs._fc.clear()


def build(light="Mivada_Standard_Light.pptx", dark="Mivada_Standard_Dark.pptx",
          outdir=HERE, basename=None, title="Mivada — the Standard"):
    if basename:
        pr, deck_html = f"{basename}-previews", f"{basename}-deck.html"
        a4l, a4p = f"{basename}-a4-landscape.html", f"{basename}-a4-portrait.html"
        docx_l, docx_d = f"{basename}.docx", f"{basename}_Dark.docx"
    else:
        pr, deck_html = "previews", "presentation.html"
        a4l, a4p = "standard-a4-landscape.html", "standard-a4-portrait.html"
        docx_l, docx_d = "Mivada_Standard.docx", "Mivada_Standard_Dark.docx"
    decks = {"light": light, "dark": dark}
    n = len(Presentation(os.path.join(outdir, light)).slides._sldIdLst)
    for theme, deck in decks.items():
        od = os.path.join(outdir, pr, theme); os.makedirs(od, exist_ok=True)
        for i in range(n):
            rs.render(os.path.join(outdir, deck), i, os.path.join(od, f"slide-{i+1:02d}.png"))
    print(f"rendered {n} slides x2 themes")
    def imgs(t): return [f"{pr}/{t}/slide-{i+1:02d}.png" for i in range(n)]

    figs = "\n".join(
        f'  <figure class="slide"><img loading="lazy" data-l="{l}" data-d="{d}" src="{l}" alt="Slide {i+1}"></figure>'
        for i, (l, d) in enumerate(zip(imgs("light"), imgs("dark"))))
    PRES = f"""<!doctype html><html lang="en"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title} · deck</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@500;700&display=swap" rel="stylesheet">
<style>
 :root{{--coral:#EA493F}} *{{box-sizing:border-box;margin:0;padding:0}}
 html,body{{background:#4d4d4d;font-family:Inter,system-ui,Arial,sans-serif}}
 .bar{{position:fixed;inset:0 0 auto 0;height:48px;display:flex;align-items:center;gap:16px;
   padding:0 18px;background:#111;color:#fff;z-index:10;font-size:13px;font-weight:600}}
 .bar b{{color:var(--coral)}} .bar .sp{{flex:1}}
 .bar button{{background:#222;color:#fff;border:1px solid #3a3a3a;border-radius:6px;
   padding:7px 14px;font:inherit;font-weight:700;cursor:pointer}} .bar button:hover{{border-color:var(--coral)}}
 .deck{{display:flex;flex-direction:column;align-items:center;gap:22px;padding:74px 0 90px}}
 .slide{{width:1280px;max-width:95vw;aspect-ratio:16/9;background:#000;
   box-shadow:0 12px 44px rgba(0,0,0,.4);scroll-snap-align:center}}
 .slide img{{width:100%;height:100%;display:block;object-fit:contain}}
 .count{{position:fixed;right:16px;bottom:14px;background:#111;color:#fff;border-radius:20px;
   padding:6px 14px;font-size:12px;font-weight:700;z-index:10}}
 @media print{{ html,body{{background:#fff}} .bar,.count{{display:none}} .deck{{gap:0;padding:0}}
   .slide{{width:100%;max-width:none;box-shadow:none;break-after:page}} }}
</style></head><body>
<div class="bar"><b>MIVADA</b> · {title} — {n} slides <span class="sp"></span>
  <button onclick="toggle()" id="tg">◑ Dark</button><button onclick="print()">⎙ Print / PDF</button></div>
<main class="deck" id="deck">
{figs}
</main>
<div class="count" id="count">1 / {n}</div>
<script>
 let dark=false;
 function toggle(){{dark=!dark;document.querySelectorAll('.slide img').forEach(im=>im.src=dark?im.dataset.d:im.dataset.l);
   document.getElementById('tg').textContent=dark?'◐ Light':'◑ Dark';document.body.style.background=dark?'#1a1a1a':'#4d4d4d';}}
 const sl=[...document.querySelectorAll('.slide')];
 function go(d){{const y=scrollY+innerHeight/2;let i=sl.findIndex(s=>s.offsetTop+s.offsetHeight/2>y);
   if(i<0)i=sl.length-1;i=Math.max(0,Math.min(sl.length-1,i+d));sl[i].scrollIntoView({{behavior:'smooth',block:'center'}});}}
 addEventListener('keydown',e=>{{if(['ArrowRight','ArrowDown','PageDown',' '].includes(e.key)){{e.preventDefault();go(1);}}
   if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){{e.preventDefault();go(-1);}} if(e.key.toLowerCase()==='d')toggle();}});
 const io=new IntersectionObserver(es=>es.forEach(en=>{{if(en.isIntersecting)
   document.getElementById('count').textContent=(sl.indexOf(en.target)+1)+' / {n}';}}),{{threshold:.6}});
 sl.forEach(s=>io.observe(s));
</script></body></html>"""
    open(os.path.join(outdir, deck_html), "w").write(PRES)

    def a4(orient):
        page = "297mm 210mm" if orient == "landscape" else "210mm 297mm"
        pw = "297mm" if orient == "landscape" else "210mm"
        ph = "210mm" if orient == "landscape" else "297mm"
        f = "\n".join(f'  <section class="pg"><img data-l="{l}" data-d="{d}" alt="Slide {i+1}" src="{l}"></section>'
                      for i, (l, d) in enumerate(zip(imgs("light"), imgs("dark"))))
        return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{title} · A4 {orient}</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@600;700&display=swap" rel="stylesheet"><style>
 *{{box-sizing:border-box;margin:0;padding:0}} html,body{{background:#4d4d4d;font-family:Inter,system-ui,Arial,sans-serif}}
 @page{{size:{page};margin:0}}
 .bar{{position:fixed;inset:0 0 auto 0;height:44px;display:flex;gap:14px;align-items:center;padding:0 16px;
   background:#111;color:#fff;z-index:9;font-size:13px;font-weight:700}}
 .bar button{{background:#222;color:#fff;border:1px solid #3a3a3a;border-radius:6px;padding:6px 12px;font:inherit;cursor:pointer}}
 .doc{{display:flex;flex-direction:column;align-items:center;gap:16px;padding:64px 0}}
 .pg{{width:{pw};height:{ph};background:#000;box-shadow:0 8px 30px rgba(0,0,0,.4);display:flex;align-items:center;justify-content:center}}
 .pg img{{width:100%;height:100%;object-fit:contain}}
 @media print{{ html,body{{background:#fff}} .bar{{display:none}} .doc{{gap:0;padding:0}} .pg{{box-shadow:none;break-after:page}} }}
</style></head><body><div class="bar"><b style="color:#EA493F">MIVADA</b> · A4 {orient} · {n} pages
  <button onclick="t()">Dark</button><button onclick="print()">Print / PDF</button></div><div class="doc">
{f}
</div><script>let d=0;function t(){{d=!d;document.querySelectorAll('.pg img').forEach(i=>i.src=d?i.dataset.d:i.dataset.l);}}</script>
</body></html>"""
    open(os.path.join(outdir, a4l), "w").write(a4("landscape"))
    open(os.path.join(outdir, a4p), "w").write(a4("portrait"))

    from PIL import Image
    from docx import Document
    from docx.shared import Mm
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    def docx(theme, outfile, bg):
        pad = os.path.join(outdir, pr, theme + "_page"); os.makedirs(pad, exist_ok=True)
        pages = []
        for i in range(n):
            im = Image.open(os.path.join(outdir, f"{pr}/{theme}/slide-{i+1:02d}.png")).convert("RGB")
            HH = round(im.width * 210 / 297)
            cv = Image.new("RGB", (im.width, HH), bg); cv.paste(im, (0, (HH - im.height) // 2))
            p = os.path.join(pad, f"p{i+1:02d}.png"); cv.save(p); pages.append(p)
        doc = Document(); s = doc.sections[0]
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Mm(297), Mm(210)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(0)
        for i, p in enumerate(pages):
            if i: doc.add_page_break()
            par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = par.paragraph_format; pf.space_before = pf.space_after = Mm(0); pf.line_spacing = 1.0
            par.add_run().add_picture(p, width=Mm(297), height=Mm(210))
        doc.save(os.path.join(outdir, outfile)); shutil.rmtree(pad, ignore_errors=True)
    docx("light", docx_l, (250, 250, 250)); docx("dark", docx_d, (0, 0, 0))
    print(f"wrote {deck_html}, {a4l}, {a4p}, {docx_l}, {docx_d}")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--light", default="Mivada_Standard_Light.pptx")
    ap.add_argument("--dark", default="Mivada_Standard_Dark.pptx")
    ap.add_argument("--outdir", default=HERE)
    ap.add_argument("--basename", default=None)
    ap.add_argument("--title", default="Mivada — the Standard")
    a = ap.parse_args()
    build(a.light, a.dark, a.outdir, a.basename, a.title)
