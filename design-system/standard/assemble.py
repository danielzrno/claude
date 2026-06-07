#!/usr/bin/env python3
"""Assemble a deck from the Standard master by picking slides, then enforce a
light/dark alternation and emit every format.

  python assemble.py --list
  python assemble.py --name PaceFarm \
      --slides cover,who-we-are,what-we-do,engagement-model,delivery-team,outcomes,contact \
      [--start dark] [--no-web]

Produces, in this folder:
  <name>.pptx              the deck — slides alternate black / white
  <name>_Light.pptx        every slide white      (theme_deck)
  <name>_Dark.pptx         every slide black       (theme_deck)
  <name>-deck.html, <name>-a4-*.html, <name>.docx, <name>_Dark.docx   (build_web, unless --no-web)
"""
import os, re, sys, shutil, argparse
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
import theme_deck, build_web

HERE = os.path.dirname(os.path.abspath(__file__))
MASTER = os.path.join(HERE, "Mivada_Standard.pptx")

# slide key -> 0-based index in the master (stable; see README "Story")
SLIDES = {
    "logo-open": 0, "cover": 1, "who-we-are": 2, "capability": 3, "trusted-by": 4,
    "what-we-do": 5, "engagement-model": 6, "engagement-models": 7, "velocity": 8,
    "timeline": 9, "data-handshake": 10, "delivery-team": 11, "payroll-testing": 12,
    "coverage": 13, "ams-itil": 14, "ams-framework": 15, "coverage-model": 16,
    "governance-model": 17, "transition": 18, "slas": 19, "commercials": 20,
    "augmentation": 21, "governance-pyramid": 22, "team-structure": 23, "outcomes": 24,
    "why-mivada": 25, "contact": 26,
}

# named presets — common decks, picked from the library in narrative order
PRESETS = {
    "company-overview": ["cover", "who-we-are", "capability", "trusted-by", "what-we-do",
                         "engagement-model", "outcomes", "why-mivada", "contact"],
    "workday-go": ["cover", "velocity", "timeline", "data-handshake", "delivery-team",
                   "payroll-testing", "commercials", "contact"],
    "ams-proposal": ["cover", "what-we-do", "ams-itil", "ams-framework", "coverage-model",
                     "governance-model", "transition", "slas", "commercials", "augmentation", "contact"],
    "full": list(SLIDES.keys()),
}

def natural_dark(slide):
    x = etree.tostring(slide.background.element).decode()
    m = re.search(r"<p:bg>.*?</p:bg>", x, re.S)
    if m and (re.search(r'srgbClr val="000000"', m.group(0)) or
              re.search(r'schemeClr val="(tx1|dk1)"', m.group(0))):
        return True
    return False

def assemble(name, keys, start=None, web=True, outdir=HERE):
    bad = [k for k in keys if k not in SLIDES]
    if bad: sys.exit(f"unknown slide keys: {bad}\nrun --list to see options")
    sel = [SLIDES[k] for k in keys]
    out = os.path.join(outdir, name + ".pptx")
    shutil.copy(MASTER, out)
    prs = Presentation(out)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep = [ids[i] for i in sel]
    for sid in list(lst): lst.remove(sid)
    for sid in keep: lst.append(sid)            # selected slides, in requested order
    for i, sid in enumerate(ids):               # drop relationships of unused slides
        if i not in sel:
            try: prs.part.drop_rel(sid.get(qn("r:id")))
            except Exception: pass
    # enforce strict light/dark alternation, starting from the opener's natural theme
    first = list(prs.slides)[0]
    cur = (start == "dark") if start in ("light", "dark") else natural_dark(first)
    for s in prs.slides:
        theme_deck.set_theme(cur); theme_deck.recolor(s); theme_deck.fix_images(s)
        cur = not cur
    prs.save(out)
    theme_deck.build(out, os.path.join(outdir, name + "_Light.pptx"), False)
    theme_deck.build(out, os.path.join(outdir, name + "_Dark.pptx"), True)
    print(f"wrote {name}.pptx (alternating) + _Light/_Dark — {len(keep)} slides")
    if web:
        build_web.build(light=name + "_Light.pptx", dark=name + "_Dark.pptx",
                        outdir=outdir, basename=name, title=name)

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--list", action="store_true", help="print the slide catalogue + presets")
    ap.add_argument("--name")
    ap.add_argument("--preset", choices=list(PRESETS), help="a named slide set")
    ap.add_argument("--slides", help="comma-separated slide keys, in order")
    ap.add_argument("--start", choices=["light", "dark"], default=None)
    ap.add_argument("--no-web", action="store_true")
    a = ap.parse_args()
    if a.list:
        print("slide keys (master order):")
        for k, i in SLIDES.items(): print(f"  {i+1:2d}  {k}")
        print("\npresets:")
        for p, ks in PRESETS.items(): print(f"  {p:18} {', '.join(ks)}")
    else:
        keys = PRESETS[a.preset] if a.preset else ([k.strip() for k in a.slides.split(",")] if a.slides else None)
        name = a.name or (a.preset.replace("-", "_").title() if a.preset else None)
        if not (name and keys): sys.exit("need --preset, or --name and --slides (or --list)")
        assemble(name, keys, a.start, not a.no_web)
