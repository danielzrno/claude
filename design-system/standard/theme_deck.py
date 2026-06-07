#!/usr/bin/env python3
"""Recolour the whole deck to a single theme (light or dark), context-aware.
   python theme_deck.py SRC.pptx {light|dark} OUT.pptx
Text colour is chosen from the luminance of the shape behind it, so the same
source works in both directions. Logos swap per theme; the governance pyramid
gets a white backing panel on dark."""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC, THEME, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
DARK = THEME == "dark"
def H(s): return RGBColor.from_string(s)
ASSET = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "logos") + os.sep
ICON  = ASSET + ("Mivada_Icon_White_RGB_L.png" if DARK else "Mivada_Icon_Melon_RGB_L.png")
WORD  = ASSET + ("Mivada_Logo_2C_OnBlack_RGB_L.png" if DARK else "Mivada_Logo_Master_RGB_L.png")

BG    = "000000" if DARK else "FAFAFA"
CARD  = "141414" if DARK else "FFFFFF"
BORDER= "3A3A3A" if DARK else "E4E4E0"
LINE  = "3A3A3A" if DARK else "AEAEAE"
TRACK = "2A2A2A" if DARK else "EFEFEA"
BAND  = "1C1C1C" if DARK else "111111"
NEUTRAL = "1A1A1A" if DARK else "F2F3EE"
ACCENT_INK = "FFFFFF" if DARK else "111111"   # thin ink accent bars

def lum(hx):
    hx = hx.upper()
    r,g,b = int(hx[0:2],16),int(hx[2:4],16),int(hx[4:6],16)
    return (0.299*r+0.587*g+0.114*b)/255

def fill_hex(sh):
    try:
        if sh.fill.type == 1: return str(sh.fill.fore_color.rgb).upper()
    except Exception: pass
    return None
def line_hex(sh):
    try:
        if sh.line.color.type is not None: return str(sh.line.color.rgb).upper()
    except Exception: pass
    return None
def set_fill(sh, hx): sh.fill.solid(); sh.fill.fore_color.rgb = H(hx)
def set_line(sh, hx): sh.line.color.rgb = H(hx)
def is_pic(sh): return sh.shape_type is not None and str(sh.shape_type).startswith("PICTURE")
def is_conn(sh): return sh._element.tag.endswith("}cxnSp") or str(sh.shape_type).startswith("LINE")

def map_fill(F, small):
    if F == "EA493F": return "EA493F"
    if F in ("FFFFFF", "141414"): return CARD
    if F == "111111": return ACCENT_INK if small else BAND
    if F == "EFEFEA": return TRACK
    if F == "F2F3EE": return NEUTRAL
    if F == "C9C9C4": return BORDER
    return F
def map_line(L):
    if L == "EA493F": return "EA493F"
    if L in ("E4E4E0", "3A3A3A"): return BORDER
    if L == "C9C9C4": return "555555" if DARK else "C9C9C4"
    if L == "AEAEAE": return LINE
    return L

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
def group_tf(g):
    xf = g._element.find(f".//{A_NS}xfrm")
    off = xf.find(f"{A_NS}off"); ext = xf.find(f"{A_NS}ext")
    cof = xf.find(f"{A_NS}chOff"); cex = xf.find(f"{A_NS}chExt")
    ox, oy = int(off.get("x")), int(off.get("y")); ew, eh = int(ext.get("cx")), int(ext.get("cy"))
    cx0, cy0 = int(cof.get("x")), int(cof.get("y"))
    cw0, ch0 = int(cex.get("cx")) or 1, int(cex.get("cy")) or 1
    sx, sy = ew / cw0, eh / ch0
    return lambda x, y, w, h: (ox + (x - cx0) * sx, oy + (y - cy0) * sy, w * sx, h * sy)
def flatten(shapes, tf):
    for sh in shapes:
        if str(sh.shape_type).startswith("GROUP"):
            if any(str(c.shape_type).startswith("PICTURE") for c in sh.shapes):
                continue   # logo wall — leave brand assets untouched
            inner = group_tf(sh)
            yield from flatten(sh.shapes, lambda x, y, w, h, o=tf, i=inner: o(*i(x, y, w, h)))
        elif sh.left is not None:
            yield sh, tf(sh.left, sh.top, sh.width, sh.height)

def recolor(slide):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = H(BG)
    flat = list(flatten(slide.shapes, lambda x, y, w, h: (x, y, w, h)))
    panels = []   # (l,t,r,b,fill) in z-order
    for sh, (ax, ay, aw, ah) in flat:
        if is_conn(sh):
            lh = line_hex(sh)
            if lh: set_line(sh, map_line(lh))
            continue
        if sh.shape_type == 1:  # autoshape
            f = fill_hex(sh)
            small = aw < Inches(0.22) or ah < Inches(0.22)
            if f:
                tgt = map_fill(f, small)
                set_fill(sh, tgt)
                if f == "111111" and not small and DARK:
                    set_line(sh, "3A3A3A")          # lift emphasis band off black
                panels.append((ax, ay, ax+aw, ay+ah, tgt))
            lh = line_hex(sh)
            if lh and not (f == "111111" and not small and DARK):
                set_line(sh, map_line(lh))
    # text pass — colour by the surface behind each textbox
    def surface(cx, cy):
        s = BG
        for (l,t,r,b,f) in panels:
            if l <= cx <= r and t <= cy <= b: s = f
        return s
    for sh, (ax, ay, aw, ah) in flat:
        if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
        cx, cy = ax + aw/2, ay + ah/2
        surf = surface(cx, cy); dark_surf = lum(surf) < 0.5
        strong_c = "FFFFFF" if dark_surf else "111111"
        muted_c  = "CFCFCF" if dark_surf else "323232"
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                try:
                    if r.font.color.type is None: continue
                    c = str(r.font.color.rgb).upper()
                except Exception:
                    continue
                if c == "EA493F":
                    r.font.color.rgb = H("FFFFFF" if surf == "EA493F" else "EA493F")
                else:
                    strong = c in ("111111", "101010", "FFFFFF", "000000")
                    r.font.color.rgb = H(strong_c if strong else muted_c)

def swap_pic(slide, sh, path):
    l,t,w,h = sh.left, sh.top, sh.width, sh.height
    sh._element.getparent().remove(sh._element)
    slide.shapes.add_picture(path, l, t, w, h)

def fix_images(slide):
    for sh in list(slide.shapes):
        if not is_pic(sh) or sh.width is None: continue
        asp = sh.width / sh.height
        if sh.width < Inches(1.0) and abs(asp - 1.69) < 0.3:
            swap_pic(slide, sh, ICON)
        elif abs(asp - 6.59) < 0.5:
            swap_pic(slide, sh, WORD)
        elif sh.width > Inches(4) and DARK:
            pad = Inches(0.18)
            panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                        sh.left-pad, sh.top-pad, sh.width+2*pad, sh.height+2*pad)
            panel.shadow.inherit = False
            panel.fill.solid(); panel.fill.fore_color.rgb = H("FFFFFF")
            panel.line.fill.background()
            sh._element.addprevious(panel._element)   # behind the image

prs = Presentation(SRC)
for s in prs.slides:
    recolor(s); fix_images(s)
prs.save(OUT)
print("saved", OUT, THEME)
