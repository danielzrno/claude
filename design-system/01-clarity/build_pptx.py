#!/usr/bin/env python3
# =============================================================================
# Mivada — Concept 01 · CLARITY  →  clarity.pptx
# Corporate / clean. White slides, ink text, coral as a disciplined accent only
# (eyebrows, one key word, a small M-chip). Black closing bookend. No shadows,
# no accent line under titles, no full-width bars, hairline rules used sparingly.
#
# Structure: a THEME dict + a Deck class of helper methods, so a new Mivada deck
# is a short list of helper calls — not hand-placed shapes.
#
# Build:  python build_pptx.py
# QA:     python /mnt/skills/public/pptx/scripts/office/soffice.py --headless \
#               --convert-to pdf clarity.pptx
# =============================================================================
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.oxml.ns import qn

# ---- Brand tokens (exact, sampled) -----------------------------------------
THEME = {
    "coral":        "EA493F",
    "coral_deep":   "C9362B",
    "black":        "000000",
    "ink":          "111111",
    "off_white":    "FAFAFA",
    "white":        "FFFFFF",
    "warm_neutral": "F2F3EE",
    "graphite":     "323232",
    "mid_grey":     "AEAEAE",
    "hairline":     "E4E4E0",
    "font":         "Inter",            # Office fallback substitutes Arial if absent
    "track_eyebrow": 0.14,              # em (applied as Pt spacing per char below)
}

EMU_IN = 914400
SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.72                           # generous side margin (board-ready)


def C(name):
    return RGBColor.from_string(THEME[name])


def _hexrgb(name):
    return RGBColor.from_string(THEME[name])


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(SLIDE_W * EMU_IN))
        self.prs.slide_height = Emu(int(SLIDE_H * EMU_IN))
        self.blank = self.prs.slide_layouts[6]

    # ---- low-level helpers --------------------------------------------------
    def _slide(self, bg="off_white"):
        s = self.prs.slides.add_slide(self.blank)
        rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                  self.prs.slide_width, self.prs.slide_height)
        rect.fill.solid()
        rect.fill.fore_color.rgb = C(bg)
        rect.line.fill.background()
        self._no_shadow(rect)
        # send the background rect to the very back
        sp = rect._element
        sp.getparent().remove(sp)
        s.shapes._spTree.insert(2, sp)
        return s

    def _no_shadow(self, shape):
        # Three steps so NO shadow renders anywhere (the Clarity brand rule):
        #  1. clear the inherited preset shadow,
        #  2. force an EMPTY <a:effectLst/> on the shape's own properties,
        #  3. strip the <p:style> block — its <a:effectRef idx="N"> otherwise
        #     re-applies the theme's drop shadow even when spPr has no effects
        #     (this is what LibreOffice was rendering).
        shape.shadow.inherit = False
        el = shape._element
        try:
            spPr = el.spPr
        except AttributeError:
            spPr = None
        if spPr is not None:
            for tag in ("a:effectLst", "a:effectDag"):
                existing = spPr.find(qn(tag))
                if existing is not None:
                    spPr.remove(existing)
            spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
        style = el.find(qn("p:style"))
        if style is not None:
            el.remove(style)

    def _text(self, slide, x, y, w, h, runs, *, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, line_spacing=None, space_after=None,
              wrap=True):
        """runs: list of (text, size, color, bold, tracking_em, italic)."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        for (txt, size, color, bold, track, *rest) in runs:
            italic = rest[0] if rest else False
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.name = THEME["font"]
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = C(color) if isinstance(color, str) and color in THEME else RGBColor.from_string(color)
            r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_AUS
            if track:
                self._set_tracking(r, size, track)
            # ensure complex/east-asian fonts also map to Inter
            self._force_font(r)
        return tb

    def _para(self, tf, runs, *, align=PP_ALIGN.LEFT, line_spacing=None,
              space_before=None, space_after=None, first=False):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        if space_before is not None:
            p.space_before = Pt(space_before)
        if space_after is not None:
            p.space_after = Pt(space_after)
        for (txt, size, color, bold, track, *rest) in runs:
            italic = rest[0] if rest else False
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.name = THEME["font"]
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = C(color) if color in THEME else RGBColor.from_string(color)
            r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_AUS
            if track:
                self._set_tracking(r, size, track)
            self._force_font(r)
        return p

    def _set_tracking(self, run, size_pt, track_em):
        # OOXML spc is in 1/100 pt; convert em→pt at the run size.
        spc = int(track_em * size_pt * 100)
        run.font._rPr.set("spc", str(spc))

    def _force_font(self, run):
        rPr = run.font._rPr
        for tag in ("a:latin", "a:cs", "a:ea"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", THEME["font"])

    def _rect(self, slide, x, y, w, h, fill=None, line=None, line_w=1.0,
              shape=MSO_SHAPE.RECTANGLE):
        sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = C(fill) if fill in THEME else RGBColor.from_string(fill)
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = C(line) if line in THEME else RGBColor.from_string(line)
            sp.line.width = Pt(line_w)
        self._no_shadow(sp)
        return sp

    def _hairline(self, slide, x, y, w, color="hairline", weight=1.0):
        ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = C(color) if color in THEME else RGBColor.from_string(color)
        ln.line.width = Pt(weight)
        self._no_shadow(ln)
        return ln

    def _vline(self, slide, x, y, h, color="hairline", weight=1.0):
        ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x), Inches(y + h))
        ln.line.color.rgb = C(color) if color in THEME else RGBColor.from_string(color)
        ln.line.width = Pt(weight)
        self._no_shadow(ln)
        return ln

    def _mchip(self, slide, x, y, size=0.62, fontcolor="white"):
        """The signature M chip: coral rounded-square + white Inter-800 M."""
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(x), Inches(y), Inches(size), Inches(size))
        chip.fill.solid()
        chip.fill.fore_color.rgb = C("coral")
        chip.line.fill.background()
        self._no_shadow(chip)
        try:
            chip.adjustments[0] = 0.22   # ~20% corner radius
        except Exception:
            pass
        tf = chip.text_frame
        tf.word_wrap = False
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "M"
        r.font.size = Pt(size * 46)      # scales with chip
        r.font.bold = True
        r.font.name = THEME["font"]
        r.font.color.rgb = C(fontcolor)
        self._set_tracking(r, size * 46, -0.04)
        self._force_font(r)
        return chip

    def _eyebrow(self, slide, x, y, w, text, color="coral"):
        return self._text(slide, x, y, w, 0.32,
                           [(text.upper(), 11, color, True, THEME["track_eyebrow"])])

    def _chrome(self, slide, num, dark=False):
        col = "AEAEAE" if not dark else "7A7A7A"
        self._text(slide, MARGIN, SLIDE_H - 0.5, 3, 0.3,
                   [("MIVADA", 9.5, col, True, 0.14)])
        self._text(slide, SLIDE_W - MARGIN - 2.5, SLIDE_H - 0.5, 2.5, 0.3,
                   [(num, 9.5, col, True, 0.14)], align=PP_ALIGN.RIGHT)

    # =========================================================================
    # SLIDE HELPERS
    # =========================================================================
    def cover(self, eyebrow, line_runs, standfirst):
        s = self._slide("off_white")
        self._mchip(s, MARGIN, 0.72, size=0.66)
        # vertically-centred block
        self._eyebrow(s, MARGIN, 2.78, 8, eyebrow)
        tf_box = self._text(s, MARGIN, 3.14, 10.6, 2.4, line_runs[0],
                            line_spacing=1.02)
        # additional headline lines
        tf = tf_box.text_frame
        for ln in line_runs[1:]:
            self._para(tf, ln, line_spacing=1.02)
        self._text(s, MARGIN, 5.42, 7.4, 1.2,
                   [(standfirst, 16.5, "graphite", False, 0)],
                   line_spacing=1.4)
        self._chrome(s, "Capability overview")
        return s

    def section_split(self, eyebrow, title_runs, standfirst_runs, num,
                      right_builder):
        """Generic left text / right content split (used by who-we-are, data)."""
        s = self._slide("off_white")
        lx = MARGIN
        self._eyebrow(s, lx, 1.02, 6.4, eyebrow)
        tb = self._text(s, lx, 1.42, 6.2, 2.4, title_runs[0], line_spacing=1.06)
        tf = tb.text_frame
        for ln in title_runs[1:]:
            self._para(tf, ln, line_spacing=1.06)
        if standfirst_runs:
            self._text(s, lx, 4.05, 5.8, 2.2, standfirst_runs,
                       line_spacing=1.42)
        right_builder(s)
        self._chrome(s, num)
        return s

    def kpis(self, eyebrow, title_runs, standfirst_runs, stats, num):
        """Who-we-are: left text + right hairline-separated stat stack."""
        def right(s):
            rx, rw = 7.55, 5.06
            top = 1.18
            row_h = 1.55
            for i, (eb, big, lab) in enumerate(stats):
                y = top + i * row_h
                if i > 0:
                    self._hairline(s, rx, y - 0.16, rw)
                self._eyebrow(s, rx, y, rw, eb)
                self._text(s, rx, y + 0.34, rw, 0.9,
                           [(big, 44, "ink", True, -0.02)])
                self._text(s, rx, y + 1.02, rw, 0.4,
                           [(lab, 12.5, "graphite", False, 0)], line_spacing=1.2)
        return self.section_split(eyebrow, title_runs, standfirst_runs, num, right)

    def pillars(self, eyebrow, title, items, num):
        """2x2 white cards, hairline border, no shadow."""
        s = self._slide("off_white")
        self._eyebrow(s, MARGIN, 0.92, 8, eyebrow)
        self._text(s, MARGIN, 1.3, 9, 0.9, [(title, 40, "ink", True, -0.02)])
        gx, gy = MARGIN, 2.62
        gw = SLIDE_W - 2 * MARGIN
        gap = 0.32
        cw = (gw - gap) / 2
        ch = 1.86
        rh_gap = 0.3
        for i, (idx, ttl, body) in enumerate(items):
            col = i % 2
            row = i // 2
            x = gx + col * (cw + gap)
            y = gy + row * (ch + rh_gap)
            card = self._rect(s, x, y, cw, ch, fill="white",
                              line="hairline", line_w=1.0,
                              shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            try:
                card.adjustments[0] = 0.035
            except Exception:
                pass
            pad = 0.34
            self._text(s, x + pad, y + 0.26, cw - 2 * pad, 0.3,
                       [(idx, 12.5, "mid_grey", True, 0.04)])
            self._text(s, x + pad, y + 0.58, cw - 2 * pad, 0.4,
                       [(ttl, 18, "ink", True, -0.01)])
            self._text(s, x + pad, y + 1.02, cw - 2 * pad, 0.8,
                       [(body, 13, "graphite", False, 0)], line_spacing=1.34)
        self._chrome(s, num)
        return s

    def steps(self, eyebrow, title, steps, num):
        """Four columns with ink numerals + coral eyebrows; top rule per col."""
        s = self._slide("off_white")
        self._eyebrow(s, MARGIN, 0.92, 9.5, eyebrow)
        self._text(s, MARGIN, 1.3, 10.5, 0.9, [(title, 38, "ink", True, -0.02)],
                   line_spacing=1.04)
        gx = MARGIN
        gw = SLIDE_W - 2 * MARGIN
        gap = 0.46
        cw = (gw - 3 * gap) / 4
        top = 3.2
        for i, (n, ttl, body) in enumerate(steps):
            x = gx + i * (cw + gap)
            # top rule (ink) — a deliberate motif above each step, not under a title
            self._rect(s, x, top, cw, 0.028, fill="ink")
            self._text(s, x, top + 0.22, cw, 0.7,
                       [(n, 38, "ink", True, -0.02)])
            self._eyebrow(s, x, top + 1.05, cw, ttl, color="coral")
            self._text(s, x, top + 1.42, cw, 1.4,
                       [(body, 13, "graphite", False, 0)], line_spacing=1.36)
        self._chrome(s, num)
        return s

    def split_caplist(self, eyebrow, title_runs, standfirst_runs, badge, caps, num):
        """Data & AI: left text + coral rule badge, right labelled list."""
        def right(s):
            rx, rw = 7.45, 5.16
            top = 1.12
            ih = 1.22
            for i, (k, v) in enumerate(caps):
                y = top + i * ih
                if i > 0:
                    self._hairline(s, rx, y - 0.1, rw)
                self._text(s, rx, y, rw, 0.34,
                           [(k, 14.5, "ink", True, -0.01)])
                self._text(s, rx, y + 0.34, rw, 0.6,
                           [(v, 12.5, "graphite", False, 0)], line_spacing=1.3)
        s = self._slide("off_white")
        lx = MARGIN
        self._eyebrow(s, lx, 1.02, 6.4, eyebrow)
        tb = self._text(s, lx, 1.42, 6.1, 2.0, title_runs[0], line_spacing=1.06)
        tf = tb.text_frame
        for ln in title_runs[1:]:
            self._para(tf, ln, line_spacing=1.06)
        if standfirst_runs:
            self._text(s, lx, 3.5, 5.7, 1.6, standfirst_runs, line_spacing=1.42)
        # coral rule + badge text (a short coral rule, not under a title)
        self._rect(s, lx, 5.5, 0.42, 0.045, fill="coral")
        self._text(s, lx + 0.6, 5.36, 5.0, 0.4,
                   [(badge, 12.5, "graphite", True, 0.02)])
        right(s)
        self._chrome(s, num)
        return s

    def stats3(self, eyebrow, title, stats, quote, by, num):
        """Three big ink numerals (coral eyebrows) + hairline + short quote."""
        s = self._slide("off_white")
        self._eyebrow(s, MARGIN, 0.92, 8, eyebrow)
        self._text(s, MARGIN, 1.3, 10, 0.9, [(title, 38, "ink", True, -0.02)])
        gx = MARGIN
        gw = SLIDE_W - 2 * MARGIN
        cw = gw / 3
        top = 2.72
        for i, (eb, big, lab) in enumerate(stats):
            x = gx + i * cw
            if i > 0:
                self._vline(s, x, top + 0.05, 1.35)
            ox = x + (0.34 if i > 0 else 0)
            self._eyebrow(s, ox, top, cw - 0.4, eb)
            self._text(s, ox, top + 0.34, cw - 0.4, 1.0,
                       [(big, 66, "ink", True, -0.03)])
            self._text(s, ox, top + 1.36, cw - 0.5, 0.4,
                       [(lab, 12.5, "graphite", False, 0)], line_spacing=1.2)
        self._hairline(s, MARGIN, 4.75, gw)
        self._text(s, MARGIN, 5.05, 9.0, 1.1,
                   [(quote, 20, "ink", False, -0.01)], line_spacing=1.34)
        self._text(s, MARGIN, 6.35, 8, 0.3,
                   [(by, 12, "mid_grey", False, 0.02)])
        self._chrome(s, num)
        return s

    def reasons(self, eyebrow, title, items, num):
        """Three hairline rows: coral numeral + title + body."""
        s = self._slide("off_white")
        self._eyebrow(s, MARGIN, 0.92, 8, eyebrow)
        self._text(s, MARGIN, 1.3, 10, 0.9, [(title, 38, "ink", True, -0.02)])
        top = 2.78
        rh = 1.32
        gw = SLIDE_W - 2 * MARGIN
        for i, (n, ttl, body) in enumerate(items):
            y = top + i * rh
            if i > 0:
                self._hairline(s, MARGIN, y - 0.12, gw)
            self._text(s, MARGIN, y, 0.9, 0.6,
                       [(n, 26, "coral", True, -0.02)])
            self._text(s, MARGIN + 0.95, y - 0.02, gw - 1.0, 0.4,
                       [(ttl, 21, "ink", True, -0.01)])
            self._text(s, MARGIN + 0.95, y + 0.42, gw - 1.0, 0.6,
                       [(body, 14.5, "graphite", False, 0)], line_spacing=1.36)
        self._chrome(s, num)
        return s

    def contact(self, eyebrow, title_runs, lines, footnote, num):
        """Black closing bookend: white text, coral accent + M chip."""
        s = self._slide("black")
        self._mchip(s, MARGIN, 0.78, size=0.62)
        # title lower-left (raised to clear the chrome footer)
        self._eyebrow(s, MARGIN, 4.18, 8, eyebrow, color="coral")
        tb = self._text(s, MARGIN, 4.56, 7.3, 1.9, title_runs[0],
                        line_spacing=1.06)
        tf = tb.text_frame
        for ln in title_runs[1:]:
            self._para(tf, ln, line_spacing=1.06)
        # contact lines right
        rx = 8.7
        ry = 1.5
        tb2 = self._text(s, rx, ry, 4.0, 2.6,
                         [(lines[0][0], 14, lines[0][1], lines[0][2], 0)],
                         line_spacing=1.5)
        tf2 = tb2.text_frame
        for (txt, col, bold) in lines[1:]:
            self._para(tf2, [(txt, 14, col, bold, 0)], line_spacing=1.5,
                       space_before=2)
        self._hairline(s, rx, 3.45, 2.7, color="323232", weight=1.0)
        self._text(s, rx, 3.62, 4.2, 0.6,
                   [(footnote, 11, "mid_grey", False, 0.03)], line_spacing=1.3)
        self._chrome(s, num, dark=True)
        return s

    def save(self, path):
        self.prs.save(path)


# =============================================================================
# BUILD THE 8 SLIDES (content + helper calls — the deliverable's real value)
# =============================================================================
def build():
    d = Deck()

    # 1 · Cover
    d.cover(
        "Capability overview · 2026",
        [
            [("Technology, ", 70, "ink", True, -0.02),
             ("human", 70, "coral", True, -0.02)],
            [("first.", 70, "ink", True, -0.02)],
        ],
        "An Australian technology consultancy that turns enterprise platforms — "
        "Workday, payroll, data and AI — into measurable human outcomes.",
    )

    # 2 · Who we are
    d.kpis(
        "Who we are",
        [
            [("A consultancy built", 38, "ink", True, -0.02)],
            [("around people, not", 38, "ink", True, -0.02)],
            [("platforms.", 38, "ink", True, -0.02)],
        ],
        [("Senior onshore leadership with cost-effective delivery — and we stay "
          "on after go-live to measure outcomes, not tickets closed.",
          16, "graphite", False, 0)],
        [
            ("Established", "2014", "Founded as LJM Infotech; Mivada today."),
            ("Specialists", "120+", "Certified consultants, not generalists."),
            ("Delivery", "AU + India", "Onshore Australia, offshore India."),
        ],
        "02",
    )

    # 3 · What we do
    d.pillars(
        "What we do",
        "Four service pillars.",
        [
            ("01", "Workday & ERP",
             "Implementation, optimisation and managed support for Workday HCM, "
             "Financials and adjacent ERP."),
            ("02", "Payroll consulting",
             "Compliant, accurate payroll across complex awards and multi-entity "
             "structures."),
            ("03", "Data & AI",
             "Lakehouse architecture, governed analytics and applied AI on your "
             "people and finance data."),
            ("04", "Intelligent automation",
             "RPA, ML and process design combined to remove low-value work."),
        ],
        "03",
    )

    # 4 · How we work
    d.steps(
        "How we work — think human first",
        "We understand the people before the technology.",
        [
            ("01", "Listen", "Understand the people and the work before the technology."),
            ("02", "Design", "Shape the platform around how teams actually operate."),
            ("03", "Deliver", "Implement in weeks-long increments, not multi-year programs."),
            ("04", "Care", "Stay on after go-live; measure outcomes, not tickets closed."),
        ],
        "04",
    )

    # 5 · Data & AI
    d.split_caplist(
        "Data & AI — spotlight",
        [
            [("So leaders decide", 38, "ink", True, -0.02)],
            [("on current data.", 38, "ink", True, -0.02)],
        ],
        [("Certified consultants design Lakehouse architectures and integrate "
          "Databricks with Workday, payroll and finance — governed, real-time, "
          "and built to be trusted.", 16, "graphite", False, 0)],
        "Databricks & Workday certified",
        [
            ("Lakehouse architecture", "Designed for governance, scale and analytics from day one."),
            ("Delta Lake pipelines", "Batch and streaming, reliable and observable."),
            ("Governed KPI store", "One trusted source for people and finance metrics."),
            ("Real-time insight", "Current data in the hands of the people deciding."),
        ],
        "05",
    )

    # 6 · Outcomes
    d.stats3(
        "Outcomes",
        "Measured where it matters.",
        [
            ("Workday", "40+", "deployments delivered."),
            ("Retention", "98%", "client retention rate."),
            ("Pace", "Weeks", "to value — not months."),
        ],
        "“They built the platform around how our teams actually work — and "
        "stayed until it stuck.”",
        "People & finance lead · enterprise client",
        "06",
    )

    # 7 · Why Mivada
    d.reasons(
        "Why Mivada",
        "Three reasons it sticks.",
        [
            ("01", "Certified specialists",
             "Workday- and Databricks-certified consultants, not generalists."),
            ("02", "Onshore + offshore",
             "Senior onshore leadership with cost-effective India delivery."),
            ("03", "People-first change",
             "Adoption built in from day one, so platforms actually get used."),
        ],
        "07",
    )

    # 8 · Contact (black bookend)
    d.contact(
        "Let's talk",
        [
            [("Let's build smarter", 36, "white", True, -0.02)],
            [("systems, with people", 36, "white", True, -0.02)],
            [("at the ", 36, "white", True, -0.02),
             ("centre.", 36, "coral", True, -0.02)],
        ],
        [
            ("mivada.com", "white", True),
            ("hello@mivada.com", "white", True),
            ("Sydney & Melbourne", "FFFFFF", False),
            ("Onshore AU & India", "FFFFFF", False),
        ],
        "Experience the Mivada difference. Figures illustrative.",
        "08",
    )

    d.save("clarity.pptx")
    print("Wrote clarity.pptx")


if __name__ == "__main__":
    build()
